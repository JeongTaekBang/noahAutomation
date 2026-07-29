#!/usr/bin/env python
"""
거래처 납기현황 조회 (Delivery Status)
=======================================

거래처가 "언제 나오냐"고 물을 때 회신할 미출고 현황표를 만듭니다.
조회 기준은 **사업자등록번호**(Business registration number), 대상은 국내(`SO_국내`)입니다.

출고 여부 판정
--------------
시트의 `Status` 컬럼을 그대로 읽지 않고 `DN_국내` 출고수량으로 **직접 계산**합니다.

`SO_국내.Status`는 수기 입력이 아니라 파워쿼리 결과를 끌어온 캐시값이라서
(`XLOOKUP(SO_ID&Line item, SO_통합[...], SO_통합[출고완료])`), 새로고침 전에 파일을 읽으면
옛 판정이 그대로 나옵니다. 실제로 이미 납품한 라인이 '미출고'로 남아 있는 경우가 있고,
그대로 회신하면 **고객에게 "아직 안 나갔습니다"라고 잘못 알리게** 됩니다.

`SO_통합[출고완료]`의 정의 자체가 DN 출고수량 기반이므로, 여기서 같은 식을 계산하면
결과는 같으면서 새로고침 여부에 의존하지 않습니다 (`DN_국내`는 직접 입력 시트라 캐시가 없음).

    출고수량 없음            → 미출고
    주문수량 - 출고수량 > 0   → 부분 출고
    출고일 없음              → 공장 출고
    그 외                    → 출고 완료   ← 회신 대상에서 제외

캐시 `Status`와 계산 결과가 어긋나면 실행 시 새로고침 안내를 띄웁니다. 회신 문서는 어차피
맞게 나가지만, 대시보드·피벗 등 같은 파일을 보는 다른 산출물도 낡았다는 신호이기 때문입니다.

분할 납기
---------
한 주문 안에서 `EXW NOAH`가 갈리면 **날짜별로 행을 나눕니다.** 실제로 흔합니다
(`SOD-2026-0264`는 32라인이 7개 날짜에 걸쳐 2027년까지). 나뉜 주문은 같은 Customer PO가
여러 줄로 보이므로, 어느 품목이 어느 날짜인지 비고에 품목명을 덧붙입니다.

사용법:
    python delivery_status.py 615-81-88675        # 사업자번호 (하이픈 유무 무관)
    python delivery_status.py 엔이에스             # 거래처명 부분일치
    python delivery_status.py --list              # 미출고가 있는 거래처 목록
    python delivery_status.py 615-81-88675 --all  # 출고완료 포함 전체
    python delivery_status.py 615-81-88675 -v     # 상세 로그
"""

from __future__ import annotations

import argparse
import datetime as dt
import logging
import sys
import unicodedata
import warnings
from pathlib import Path

import pandas as pd

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

from po_generator.cli_common import generate_output_filename
from po_generator.config import (
    DS_OUTPUT_DIR,
    DN_DOMESTIC_SHEET,
    NOAH_SO_PO_DN_FILE,
    SO_DOMESTIC_SHEET,
)
from po_generator.logging_config import setup_logging
from po_generator.utils import normalize_biz_no, normalize_line_item

logger = logging.getLogger(__name__)


# 회신 대상 출고상태 — '출고 완료'만 빠진다
PENDING_STATUSES: tuple[str, ...] = ('미출고', '부분 출고', '공장 출고')

# 회신에서 통째로 빼는 주문 (캐시 여부와 무관하게 시트 Status로만 판단)
# 취소/보류는 DN이 영영 안 생기므로 수량으로는 구분할 수 없다.
EXCLUDED_SO_STATUSES: frozenset[str] = frozenset({'Cancelled', 'Hold'})

# EXW NOAH가 비어 있을 때 고객 회신에 찍을 문구
DATE_TBD_LABEL = '처리중'

# 한 주문의 납기가 갈려 행이 나뉠 때, 비고에 덧붙일 품목명의 최대 길이 (넘으면 '외 N종')
# 개수가 아니라 길이로 자른다 — 품목명이 'MA02/0.75kW/43RPM/ON-OFF/SBWG-04-1SM'처럼
# 긴 거래처(피엠에스)가 있어서 "3개까지"로는 회신 표가 깨진다.
SPLIT_ITEM_NOTE_MAX_CHARS = 60

# 사업자등록번호 자릿수 — 조회어가 번호인지 이름인지 가르는 기준
BIZ_NO_MIN_DIGITS = 8


# === 값 정규화 =============================================================

def as_date(value) -> pd.Timestamp | None:
    """날짜 셀 → Timestamp (빈 값이면 None)

    `EXW NOAH`는 빈 칸이 NaN이 아니라 `datetime.time(0, 0)`으로 읽힙니다
    (셀에 0 값이 들어 있어 컬럼 dtype이 object로 떨어진다).
    NaN만 검사하면 미정 납기가 1900-01-00 같은 날짜로 둔갑하므로 타입까지 본다.
    """
    if value is None or value is pd.NaT:
        return None
    if isinstance(value, dt.time):  # 빈 칸이 시간 0으로 읽힌 경우
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    if isinstance(value, str) and not value.strip():
        return None
    try:
        ts = pd.Timestamp(value)
    except (ValueError, TypeError):
        return None
    if pd.isna(ts):
        return None
    return ts


def format_date(value) -> str:
    """날짜 → 'YYYY-MM-DD' (빈 값이면 '')"""
    ts = as_date(value)
    return ts.strftime('%Y-%m-%d') if ts is not None else ''


def _clean_text(value) -> str:
    """셀 값 → 표시용 문자열 (NaN/None은 빈 문자열)"""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ''
    text = str(value).strip()
    return '' if text.lower() == 'nan' else text


def _display_width(text: str) -> int:
    """콘솔 표시 폭 — 한글/한자는 2칸을 차지한다"""
    return sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in text)


def _pad(text: str, width: int) -> str:
    """표시 폭 기준 왼쪽 정렬 패딩 (넘치면 '…'로 자른다)

    f-string의 `:<32`는 글자 수로 세기 때문에 거래처명·Remarks가 한글이면 열이 어긋난다.
    """
    if _display_width(text) > width:
        out = ''
        for ch in text:
            if _display_width(out) + _display_width(ch) > width - 1:
                break
            out += ch
        text = out + '…'
    return text + ' ' * max(0, width - _display_width(text))


# === 데이터 로드 ===========================================================

def load_sheets() -> tuple[pd.DataFrame, pd.DataFrame]:
    """SO_국내 / DN_국내 로드

    Returns:
        (SO DataFrame, DN DataFrame)

    Raises:
        FileNotFoundError: 데이터 파일이 없는 경우
    """
    if not NOAH_SO_PO_DN_FILE.exists():
        raise FileNotFoundError(f"소스 파일을 찾을 수 없습니다: {NOAH_SO_PO_DN_FILE}")

    logger.info("데이터 로드: %s", NOAH_SO_PO_DN_FILE.name)
    with pd.ExcelFile(NOAH_SO_PO_DN_FILE) as xl:
        so = pd.read_excel(xl, sheet_name=SO_DOMESTIC_SHEET)
        dn = pd.read_excel(xl, sheet_name=DN_DOMESTIC_SHEET)

    so = so[so['SO_ID'].notna()].copy()
    dn = dn[dn['DN_ID'].notna()].copy()
    logger.info("SO_국내 %d행 / DN_국내 %d행", len(so), len(dn))
    return so, dn


def aggregate_dn(dn: pd.DataFrame) -> pd.DataFrame:
    """DN_국내를 SO_ID + Line item 단위 출고수량/최종출고일로 집계

    한 SO 라인을 여러 번에 나눠 출고(분할 납품)하면 DN 행이 여러 개이므로 합산한다.
    """
    columns = ['_so_id', '_line', '출고수량', '_last_ship', '_dn_match']
    if dn.empty:
        # dtype을 명시한다 — 빈 프레임을 columns만으로 만들면 전부 object가 되고,
        # merge 뒤 fillna에서 pandas 다운캐스팅 경고가 난다.
        return pd.DataFrame({
            '_so_id': pd.Series(dtype='object'),
            '_line': pd.Series(dtype='object'),
            '출고수량': pd.Series(dtype='float64'),
            '_last_ship': pd.Series(dtype='datetime64[ns]'),
            '_dn_match': pd.Series(dtype='object'),
        })

    work = pd.DataFrame({
        '_so_id': dn['SO_ID'].map(_clean_text),
        '_line': dn['Line item'].map(normalize_line_item),
        '출고수량': pd.to_numeric(dn['Qty'], errors='coerce').fillna(0.0),
        # as_date로 한 번 거른 뒤 datetime64로 — 빈 칸이 time(0,0)으로 읽히는 시트가 섞여도
        # groupby('max')가 object dtype에서 깨지지 않게 한다.
        '_ship_date': pd.to_datetime(dn['출고일'].map(as_date), errors='coerce'),
    })
    grouped = work.groupby(['_so_id', '_line'], as_index=False).agg(
        출고수량=('출고수량', 'sum'),
        _last_ship=('_ship_date', 'max'),
    )
    # DN 행의 '존재' 자체를 표시 — 파워쿼리의 `[출고수량] = null` 판정과 맞추기 위해서다.
    # 수량 0짜리 DN 행이 있으면 '미출고'가 아니라 '부분 출고'가 맞다.
    grouped['_dn_match'] = True
    return grouped[columns]


def attach_ship_status(so: pd.DataFrame, dn: pd.DataFrame) -> pd.DataFrame:
    """SO 라인에 출고수량·미출고수량·출고상태를 붙인다

    출고상태 판정은 파워쿼리 `SO_통합[출고완료]`와 같은 식이다 (모듈 docstring 참조).
    """
    dn_agg = aggregate_dn(dn)

    work = so.copy()
    work['_so_id'] = work['SO_ID'].map(_clean_text)
    work['_line'] = work['Line item'].map(normalize_line_item)
    work['_biz'] = work['Business registration number'].map(normalize_biz_no)
    work['_주문수량'] = pd.to_numeric(work['Item qty'], errors='coerce').fillna(0.0)
    work['_단가'] = pd.to_numeric(work['Sales Unit Price'], errors='coerce').fillna(0.0)

    # SO 라인이 (SO_ID, Line item)으로 유일하지 않으면 merge가 fan-out되어 수량이 부풀려진다.
    dup = int(work.duplicated(subset=['_so_id', '_line']).sum())
    if dup:
        logger.warning(
            "%s: (SO_ID, Line item) 중복 %d건 — 출고수량이 중복 집계될 수 있습니다",
            SO_DOMESTIC_SHEET, dup,
        )

    work = work.merge(dn_agg, on=['_so_id', '_line'], how='left')
    work['출고수량'] = pd.to_numeric(work['출고수량'], errors='coerce').fillna(0.0)
    # merge 결과의 _dn_match는 True 아니면 NaN이라 notna()가 곧 '매칭됨'이다
    # (object 컬럼에 fillna(False)를 쓰면 pandas 다운캐스팅 경고가 난다)
    work['_has_dn'] = work['_dn_match'].notna()
    work['_미출고수량'] = (work['_주문수량'] - work['출고수량']).clip(lower=0.0)

    def _status(row) -> str:
        if not row['_has_dn']:
            return '미출고'
        if row['_주문수량'] - row['출고수량'] > 0.001:
            return '부분 출고'
        if pd.isna(row['_last_ship']):
            return '공장 출고'
        return '출고 완료'

    # 빈 프레임에 apply(axis=1)을 걸면 Series가 아니라 DataFrame이 돌아와 대입이 깨진다
    work['_출고상태'] = (
        work.apply(_status, axis=1) if len(work)
        else pd.Series(dtype='object', index=work.index)
    )
    work['_미출고금액'] = work['_단가'] * work['_미출고수량']
    return work


def cached_status(df: pd.DataFrame) -> pd.Series:
    """시트의 `Status` 컬럼을 문자열 Series로 (컬럼이 없으면 전부 빈 문자열)

    Status는 파워쿼리 결과의 캐시라 언제든 통째로 빠질 수 있다 (수식 오류, 시트 개편).
    없다고 조회가 멈추면 안 되므로 '판정 정보 없음'으로 취급한다.
    """
    if 'Status' not in df.columns:
        return pd.Series('', index=df.index, dtype=object)
    return df['Status'].map(_clean_text)


def count_stale_status(work: pd.DataFrame) -> int:
    """시트의 캐시 `Status`와 계산된 출고상태가 어긋나는 행 수

    파워쿼리 새로고침이 밀렸는지 알려주는 신호. 취소/보류 주문은 애초에 판정 대상이
    아니므로 제외하고, 수식이 빈 문자열을 반환한 행(NaN)도 '아직 계산 안 됨'이라 제외한다.
    """
    cached = cached_status(work)
    comparable = (cached != '') & (~cached.isin(EXCLUDED_SO_STATUSES))
    return int((comparable & (cached != work['_출고상태'])).sum())


# === 거래처 조회 ===========================================================

def customer_directory(work: pd.DataFrame) -> pd.DataFrame:
    """사업자번호 → 거래처명 목록 (SO에 등장하는 거래처)

    같은 사업자번호에 표기가 여러 개면 가장 많이 쓰인 이름을 대표로 삼는다.
    """
    rows = work[work['_biz'] != ''].copy()
    rows['_name'] = rows['Customer name'].map(_clean_text)
    if rows.empty:
        return pd.DataFrame(columns=['_biz', '_name'])
    top = (
        rows.groupby(['_biz', '_name'], as_index=False)
            .size()
            .sort_values(['_biz', 'size'], ascending=[True, False])
            .drop_duplicates(subset='_biz', keep='first')
    )
    return top[['_biz', '_name']].reset_index(drop=True)


def resolve_customer(work: pd.DataFrame, query: str) -> tuple[list[tuple[str, str]], str]:
    """조회어(사업자번호 또는 거래처명) → 후보 [(사업자번호, 거래처명)]

    Returns:
        (후보 목록, 조회 방식 설명)
    """
    directory = customer_directory(work)
    digits = normalize_biz_no(query)

    if len(digits) >= BIZ_NO_MIN_DIGITS:
        hits = directory[directory['_biz'] == digits]
        return list(hits.itertuples(index=False, name=None)), '사업자번호로'

    keyword = query.strip().lower()
    hits = directory[directory['_name'].str.lower().str.contains(keyword, regex=False, na=False)]
    return list(hits.itertuples(index=False, name=None)), '거래처명으로'


def format_biz_no(digits: str) -> str:
    """정규화된 사업자번호 → 'NNN-NN-NNNNN' 표기 (10자리가 아니면 원문 유지)"""
    if len(digits) == 10:
        return f"{digits[:3]}-{digits[3:5]}-{digits[5:]}"
    return digits


# === 집계 =================================================================

def select_rows(work: pd.DataFrame, biz_no: str, include_shipped: bool) -> pd.DataFrame:
    """회신 대상 SO 라인 선별"""
    rows = work[work['_biz'] == biz_no].copy()
    rows = rows[~cached_status(rows).isin(EXCLUDED_SO_STATUSES)]

    if not include_shipped:
        rows = rows[rows['_출고상태'].isin(PENDING_STATUSES)]
    return rows


def item_note(group: pd.DataFrame) -> str:
    """행에 묶인 품목명 — 비고에 붙여 "어느 품목이 이 날짜인지" 알려준다

    품목이 많으면 회신 표가 읽히지 않으므로 길이가 찰 때까지만 쓰고 나머지는 종수로 접는다.
    첫 품목은 아무리 길어도 반드시 남긴다 (전부 '외 N종'만 나오면 단서가 없다).
    """
    names: list[str] = []
    for name in group['Item name'].map(_clean_text):
        if name and name not in names:
            names.append(name)
    if not names:
        return ''

    shown: list[str] = []
    used = 0
    for name in names:
        if shown and used + len(name) + 2 > SPLIT_ITEM_NOTE_MAX_CHARS:
            break
        shown.append(name)
        used += len(name) + 2

    text = ', '.join(shown)
    if len(names) > len(shown):
        text += f" 외 {len(names) - len(shown)}종"
    return f"품목: {text}"


def build_summary(rows: pd.DataFrame) -> pd.DataFrame:
    """주문 × 공장 출고일 단위 요약 — 고객에게 보내는 표

    기본 단위는 이미 메일로 회신하던 표와 같은 주문(SO_ID) 1행이고, 여기에 Sales 금액과
    PO receipt date를 더했다.

    다만 **한 주문 안에서 `EXW NOAH`가 갈리면 날짜별로 행을 나눈다.** 분할 납기가 실제로
    있다 (세진밸브 `SOD-2026-0264`는 32라인이 7개 날짜에 걸쳐 2027년까지). 대표 날짜 하나로
    뭉개면 나머지 납기가 통째로 사라지거나 "그날 다 나온다"는 틀린 약속이 된다.
    날짜가 하나뿐인 주문(139건 중 129건)은 예전 그대로 1행이다.

    나뉜 주문은 같은 Customer PO가 여러 줄로 보이므로 어느 행이 무엇인지 알려줘야 하는데,
    **비고만으로 구분이 안 될 때만** 품목명을 덧붙인다. 실측 10건 중 8건은 비고가 이미
    호선별로 갈려 있어(세진밸브 `H2734`/`H2735`…) 품목이 잡음일 뿐이고, 나머지 2건만
    단서가 필요하다 — 피엠에스는 두 행 모두 `묘도 GS`, 티에스엔텍은 비고가 아예 비어 있다.
    """
    if rows.empty:
        return pd.DataFrame(columns=[
            'Customer PO', 'Remarks', '수량', 'NOAH 공장 출고일', 'Sales 금액', 'PO receipt date',
        ])

    work = rows.copy()
    # 리스트를 그대로 컬럼에 넣으면 안 된다 — None이 섞이면 pandas가 datetime64로 캐스팅해
    # None이 NaT가 되고, `is not None`을 통과한 NaT가 strftime에서 터진다.
    # 날짜 계산은 컬럼에 담기 전에 끝내고, 컬럼에는 이미 만든 문자열/정렬키만 넣는다.
    exw_dates = [as_date(v) for v in work['EXW NOAH']]
    work['_exw_key'] = [
        d.strftime('%Y-%m-%d') if d is not None else DATE_TBD_LABEL for d in exw_dates
    ]
    work['_sort_exw'] = [d if d is not None else pd.Timestamp.max for d in exw_dates]

    records = []
    for _so_id, so_group in work.groupby('_so_id', sort=False):
        date_groups = list(so_group.groupby('_exw_key', sort=False))
        # Remarks가 라인마다 갈리는 주문이 있다 — 행마다 첫 비어있지 않은 값을 쓴다
        base_remarks = [
            next((r for r in g['Remarks'].map(_clean_text) if r), '')
            for _, g in date_groups
        ]
        # 품목은 "비고만으로 어느 행인지 모를 때"만 붙인다.
        # 세진밸브(SOD-2026-0264)처럼 비고가 이미 호선별로 갈려 있으면 품목은 잡음일 뿐이고,
        # 피엠에스(두 행 모두 '묘도 GS')·티에스엔텍(비고 없음)처럼 겹치면 품목이 유일한 단서다.
        need_note = len(date_groups) > 1 and (
            len(set(base_remarks)) < len(base_remarks) or any(not r for r in base_remarks)
        )

        for (exw_key, group), remarks in zip(date_groups, base_remarks):
            if need_note:
                note = item_note(group)
                if note:
                    remarks = f"{remarks} ({note})" if remarks else note
            first = group.iloc[0]
            records.append({
                'Customer PO': _clean_text(first['Customer PO']),
                'Remarks': remarks,
                '수량': float(group['_미출고수량'].sum()),
                'NOAH 공장 출고일': exw_key,
                'Sales 금액': float(group['_미출고금액'].sum()),
                'PO receipt date': format_date(first['PO receipt date']),
                '_sort_exw': first['_sort_exw'],
            })

    summary = pd.DataFrame(records)
    # 납기 확정 건을 앞에, '처리중'을 뒤에 — 회신 받는 쪽이 먼저 보고 싶은 순서
    summary = summary.sort_values(
        ['_sort_exw', 'PO receipt date', 'Customer PO'],
    ).reset_index(drop=True)
    return summary.drop(columns=['_sort_exw'])


def build_detail(rows: pd.DataFrame) -> pd.DataFrame:
    """SO 라인 단위 상세 — 내부 확인용"""
    columns = [
        'SO_ID', 'Line item', 'Customer PO', 'Remarks', 'Item name',
        '주문수량', '출고수량', '미출고수량', 'Sales Unit Price', '미출고금액',
        'PO receipt date', 'EXW NOAH', 'Expected delivery date', '출고상태',
    ]
    if rows.empty:
        return pd.DataFrame(columns=columns)

    detail = pd.DataFrame({
        'SO_ID': rows['_so_id'],
        'Line item': rows['_line'],
        'Customer PO': rows['Customer PO'].map(_clean_text),
        'Remarks': rows['Remarks'].map(_clean_text),
        'Item name': rows['Item name'].map(_clean_text),
        '주문수량': rows['_주문수량'],
        '출고수량': rows['출고수량'],
        '미출고수량': rows['_미출고수량'],
        'Sales Unit Price': rows['_단가'],
        '미출고금액': rows['_미출고금액'],
        'PO receipt date': rows['PO receipt date'].map(format_date),
        'EXW NOAH': rows['EXW NOAH'].map(format_date),
        'Expected delivery date': rows['Expected delivery date'].map(format_date),
        '출고상태': rows['_출고상태'],
    })
    detail['_line_sort'] = pd.to_numeric(detail['Line item'], errors='coerce')
    detail = detail.sort_values(['SO_ID', '_line_sort']).reset_index(drop=True)
    return detail.drop(columns=['_line_sort'])[columns]


def build_customer_list(work: pd.DataFrame) -> pd.DataFrame:
    """미출고 잔량이 있는 거래처 목록 (--list)"""
    pending = work[
        (~cached_status(work).isin(EXCLUDED_SO_STATUSES))
        & (work['_출고상태'].isin(PENDING_STATUSES))
        & (work['_biz'] != '')
    ].copy()
    if pending.empty:
        return pd.DataFrame(columns=['사업자번호', '거래처명', '주문건수', '미출고수량', '미출고금액'])

    directory = customer_directory(work).set_index('_biz')['_name'].to_dict()
    agg = pending.groupby('_biz').agg(
        주문건수=('_so_id', 'nunique'),
        미출고수량=('_미출고수량', 'sum'),
        미출고금액=('_미출고금액', 'sum'),
    ).reset_index()
    agg['거래처명'] = agg['_biz'].map(directory).fillna('')
    agg['사업자번호'] = agg['_biz'].map(format_biz_no)
    agg = agg.sort_values('미출고금액', ascending=False).reset_index(drop=True)
    return agg[['사업자번호', '거래처명', '주문건수', '미출고수량', '미출고금액']]


# === 출력 =================================================================

def write_output(
    summary: pd.DataFrame,
    detail: pd.DataFrame,
    output_file: Path,
    customer_name: str,
    biz_no: str,
) -> None:
    """납기현황 xlsx 출력 (납기현황 / 상세 2시트)"""
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    output_file.parent.mkdir(parents=True, exist_ok=True)

    def _add_table(writer, sheet_name: str, display_name: str) -> None:
        ws = writer.sheets[sheet_name]
        if ws.max_row < 2:
            return
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        tbl = Table(displayName=display_name, ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tbl)

    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        summary.to_excel(writer, sheet_name='납기현황', index=False)
        _add_table(writer, '납기현황', '납기현황')

        detail.to_excel(writer, sheet_name='상세', index=False)
        _add_table(writer, '상세', '상세')

        # 숫자 서식 + 열 너비
        widths = {
            # B(비고)는 분할 납기 때 '(품목: ...)'가 덧붙어 길어진다
            '납기현황': {'A': 22, 'B': 52, 'C': 8, 'D': 18, 'E': 16, 'F': 16},
            '상세': {
                'A': 16, 'B': 8, 'C': 22, 'D': 34, 'E': 40, 'F': 10, 'G': 10,
                'H': 11, 'I': 14, 'J': 14, 'K': 16, 'L': 14, 'M': 20, 'N': 11,
            },
        }
        money_cols = {'납기현황': ('E',), '상세': ('I', 'J')}
        for sheet, widths_map in widths.items():
            ws = writer.sheets[sheet]
            for col, width in widths_map.items():
                ws.column_dimensions[col].width = width
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal='center', vertical='center')
            for col in money_cols[sheet]:
                for row in range(2, ws.max_row + 1):
                    ws[f"{col}{row}"].number_format = '#,##0'

        # 거래처/기준일은 인쇄 머리글로 — 표 자체는 A1부터 시작해야 필터/복사가 깔끔하다
        head = writer.sheets['납기현황']
        head.oddHeader.left.text = f"{customer_name} ({format_biz_no(biz_no)})"
        head.oddHeader.right.text = f"기준일 {dt.date.today():%Y-%m-%d}"


def print_summary(summary: pd.DataFrame, customer_name: str, biz_no: str) -> None:
    """콘솔 요약 출력"""
    print()
    print(f"납기현황 — {customer_name} ({format_biz_no(biz_no)})")
    print("=" * 88)

    if summary.empty:
        print("  미출고 건이 없습니다.")
        return

    print(f"  {_pad('Customer PO', 22)} {_pad('Remarks', 32)} {'수량':>6} {_pad('공장 출고일', 12)} {'Sales 금액':>14}")
    print("  " + "-" * 86)
    for _, row in summary.iterrows():
        print(
            f"  {_pad(row['Customer PO'], 22)} {_pad(row['Remarks'], 32)} {row['수량']:>6,.0f} "
            f"{_pad(row['NOAH 공장 출고일'], 12)} {row['Sales 금액']:>14,.0f}"
        )
    print("  " + "-" * 86)
    print(
        f"  {_pad('합계', 22)} {_pad('', 32)} {summary['수량'].sum():>6,.0f} "
        f"{_pad('', 12)} {summary['Sales 금액'].sum():>14,.0f}"
    )
    tbd = int((summary['NOAH 공장 출고일'] == DATE_TBD_LABEL).sum())
    if tbd:
        print(f"\n  공장 출고일 미정({DATE_TBD_LABEL}): {tbd}건 — EXW NOAH가 비어 있습니다.")


def print_customer_list(listing: pd.DataFrame) -> None:
    """--list 출력"""
    print()
    print("미출고 잔량이 있는 거래처")
    print("=" * 76)
    if listing.empty:
        print("  없습니다.")
        return
    print(f"  {_pad('사업자번호', 16)} {_pad('거래처명', 28)} {'주문':>5} {'수량':>7} {'미출고금액':>15}")
    print("  " + "-" * 74)
    for _, row in listing.iterrows():
        print(
            f"  {_pad(row['사업자번호'], 16)} {_pad(row['거래처명'], 28)} {row['주문건수']:>5} "
            f"{row['미출고수량']:>7,.0f} {row['미출고금액']:>15,.0f}"
        )
    print("  " + "-" * 74)
    print(f"  거래처 {len(listing)}곳 / 미출고금액 합계 {listing['미출고금액'].sum():,.0f}")


def warn_if_stale(stale_count: int) -> None:
    """캐시 Status가 계산 결과와 어긋나면 새로고침 안내"""
    if stale_count <= 0:
        return
    print()
    print(f"[주의] SO_국내 Status가 실제 출고와 다른 행 {stale_count}건 — 파워쿼리 새로고침이 밀렸습니다.")
    print("       이 회신 문서는 DN_국내 출고수량으로 직접 계산하므로 영향이 없지만,")
    print("       같은 파일을 보는 대시보드·피벗은 낡은 값을 씁니다. 엑셀에서 [모두 새로 고침] 하세요.")


# === CLI ==================================================================

def create_argument_parser() -> argparse.ArgumentParser:
    """CLI 인자 파서 생성"""
    parser = argparse.ArgumentParser(
        prog='delivery_status',
        description='거래처 납기현황 조회 — 사업자등록번호 기준 미출고 현황 xlsx 생성',
        epilog=(
            '예시:\n'
            '  python delivery_status.py 615-81-88675\n'
            '  python delivery_status.py 엔이에스\n'
            '  python delivery_status.py --list'
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument(
        'customer',
        nargs='?',
        help='사업자등록번호 (하이픈 유무 무관) 또는 거래처명 일부',
    )
    parser.add_argument(
        '-l', '--list',
        action='store_true',
        dest='show_list',
        help='미출고 잔량이 있는 거래처 목록만 출력',
    )
    parser.add_argument(
        '-a', '--all',
        action='store_true',
        help='출고완료 건까지 포함 (기본은 미출고만)',
    )
    parser.add_argument(
        '-v', '--verbose',
        action='store_true',
        help='상세 로그 출력',
    )
    return parser


def main() -> int:
    """메인 함수"""
    parser = create_argument_parser()
    args = parser.parse_args()
    setup_logging(verbose=args.verbose)

    if not args.show_list and not args.customer:
        parser.print_help()
        print("\n[오류] 사업자등록번호 또는 거래처명을 입력하세요 (목록은 --list).")
        return 1

    # 1. 데이터 로드 + 출고상태 계산
    try:
        so, dn = load_sheets()
    except FileNotFoundError as e:
        print(f"[오류] {e}")
        return 1
    except PermissionError:
        print(f"[오류] 파일이 열려 있습니다. 엑셀을 닫고 다시 실행하세요: {NOAH_SO_PO_DN_FILE.name}")
        return 1

    missing = [c for c in ('SO_ID', 'Line item', 'Business registration number') if c not in so.columns]
    if missing:
        print(f"[오류] {SO_DOMESTIC_SHEET} 시트에 필요한 컬럼이 없습니다: {missing}")
        return 1

    work = attach_ship_status(so, dn)
    stale = count_stale_status(work)

    # 2. 목록 모드
    if args.show_list:
        print_customer_list(build_customer_list(work))
        warn_if_stale(stale)
        return 0

    # 3. 거래처 확정
    candidates, lookup_kind = resolve_customer(work, args.customer)
    if not candidates:
        print(f"[오류] {lookup_kind} '{args.customer}'에 해당하는 거래처를 찾을 수 없습니다.")
        print("       `python delivery_status.py --list`로 거래처 목록을 확인하세요.")
        return 1
    if len(candidates) > 1:
        print(f"[오류] '{args.customer}'에 해당하는 거래처가 {len(candidates)}곳입니다. 사업자번호로 지정하세요:")
        for biz, name in candidates:
            print(f"  - {format_biz_no(biz)}  {name}")
        return 1

    biz_no, customer_name = candidates[0]

    # 4. 집계
    rows = select_rows(work, biz_no, include_shipped=args.all)
    summary = build_summary(rows)
    detail = build_detail(rows)

    if summary.empty:
        scope = '주문' if args.all else '미출고 건'
        print(f"\n{customer_name} ({format_biz_no(biz_no)}) — 해당하는 {scope}이 없습니다.")
        warn_if_stale(stale)
        return 0

    # 5. 출력
    output_file = generate_output_filename('납기현황', customer_name, format_biz_no(biz_no), DS_OUTPUT_DIR)
    try:
        write_output(summary, detail, output_file, customer_name, biz_no)
    except PermissionError:
        print(f"[오류] 출력 파일이 열려 있습니다. 닫고 다시 실행하세요: {output_file.name}")
        return 1

    print_summary(summary, customer_name, biz_no)
    print(f"\n출력: {output_file}")
    warn_if_stale(stale)
    return 0


if __name__ == "__main__":
    sys.exit(main())
