#!/usr/bin/env python
"""
DN 출고기록 자동 입력 (Delivery Note Recorder)
==============================================

공장 출고리스트(`po_reconciliation/{year}/{period}/2026리스트_RCK_{period}.xlsx`)를 읽어
`NOAH_SO_PO_DN.xlsx`의 `DN_국내` 표에 출고 내역을 추가합니다.

무엇을 어디서 가져오는가
------------------------
    출고 이벤트   출고리스트 Delivery 시트에서 `SO_ID`가 있는 행
                 (N/A는 서비스 출고라 DN 대상이 아니다 — P08 기준 48행 중 18행)
    출고일        출고리스트 `납품완료`
    라인/수량     `PO_국내` Status가 `Invoiced P{XX}`인 라인.
                 그 SO의 PO가 전부 정리됐으면(= 남은 게 없으면) `SO_국내` 전 라인의
                 **잔량**을 쓴다 — PO는 부속을 1라인에 합쳐 적기 때문에
                 (`SA09X-MA + ADAPTER`) PO만 보면 부속 라인이 통째로 빠진다.
    DN_ID         출고리스트 1행당 1개, 기존 최대값 다음부터 연번

자동 입력 대상은 **자기검증을 통과한 건만**입니다:
    (a) Σ(PO Invoiced P{XX}의 Total ICO) == 출고리스트 계산서금액
    (b) 출고리스트에 같은 SO_ID가 여러 날짜로 있지 않을 것 (라인을 날짜별로 나눌 근거가 없음)
걸린 건은 워크북에 쓰지 않고 '확인 필요' 목록으로만 냅니다.
2026-03~08 561건 리플레이 기준 자동 입력 549건 중 548건 정확(99.8%), 헛경보 0건.

세금계산서 발행일은 출고일과 같게 채우되, 기존 DN에서 유도한 **월합 세금계산서 거래처**는
비우고 Remarks를 물려줍니다 (그 거래처는 출고 시점에 계산서를 끊지 않는다).

같은 기간을 여러 번 돌려도 안전합니다 — 이미 기록된 `(SO_ID, 출고일)`은 건너뜁니다.

사용법:
    python create_dn.py P08                # 미리보기 → y/N → 워크북에 추가
    python create_dn.py P08 --dry-run      # 미리보기 파일만 (워크북 안 건드림)
    python create_dn.py P08 --yes          # 확인 없이 추가 (배치용)
    python create_dn.py P08 --no-tax-date  # 세금계산서 발행일 전부 공란
    python create_dn.py P08 -v             # 상세 로그
"""

from __future__ import annotations

import argparse
import logging
import sys
import warnings
from pathlib import Path

import pandas as pd

warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

from po_generator.cli_common import generate_output_filename
from po_generator.config import (
    BASE_DIR,
    DN_BACKUP_DIR,
    DN_DOMESTIC_SHEET,
    DN_OUTPUT_DIR,
    MSG_ERROR,
    MSG_NOTICE,
    NOAH_SO_PO_DN_FILE,
    PO_DOMESTIC_SHEET,
    SO_DOMESTIC_SHEET,
)
from po_generator.dn_recorder import (
    Plan,
    build_plan,
    find_delivery_file,
    load_delivery,
    load_source_frames,
)
from po_generator.logging_config import setup_logging
from po_generator.mail_cli import confirm

logger = logging.getLogger(__name__)

RECON_DIR = BASE_DIR / "po_reconciliation"

# 미리보기 '추가분' 시트 — DN_국내와 같은 순서로 보여 준다 (수식 열은 '(수식)' 표시)
PREVIEW_COLUMNS: tuple[str, ...] = (
    'DN_ID', 'SO_ID', 'Line item', 'Item(참고)', 'Qty', 'Currency',
    '출고일', '세금계산서 발행일', 'Remarks', 'Customer(참고)',
)


def build_preview(plan: Plan) -> pd.DataFrame:
    """추가할 행을 사람이 눈으로 훑을 수 있는 표로"""
    return pd.DataFrame([{
        'DN_ID': line.dn_id,
        'SO_ID': line.so_id,
        'Line item': line.line_item,
        'Item(참고)': line.item_name,
        'Qty': line.qty,
        'Currency': line.currency,
        '출고일': line.ship_date,
        '세금계산서 발행일': line.tax_date,
        'Remarks': line.remarks or '',
        'Customer(참고)': line.customer_name,
    } for line in plan.lines], columns=list(PREVIEW_COLUMNS))


def build_review_table(plan: Plan) -> pd.DataFrame:
    rows = [{
        '구분': '확인 필요',
        'SO_ID': r.so_id,
        '출고일': r.ship_date,
        'Customer': r.customer,
        '사유': r.reason,
        '내용': r.detail,
    } for r in plan.reviews]
    rows += [{
        '구분': '건너뜀',
        'SO_ID': r.so_id,
        '출고일': r.ship_date,
        'Customer': r.customer,
        '사유': r.reason,
        '내용': r.detail,
    } for r in plan.skipped]
    return pd.DataFrame(rows, columns=['구분', 'SO_ID', '출고일', 'Customer', '사유', '내용'])


def write_preview(plan: Plan, output_dir: Path) -> Path:
    """미리보기 xlsx (추가분 / 확인필요 2시트)"""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    output_dir.mkdir(parents=True, exist_ok=True)
    # 출고일 범위를 파일명에 넣는다 — 월중에 여러 번 돌리므로 어느 실행이 어느 날짜분인지
    # 파일명만 보고 구분돼야 한다
    dates = sorted({l.ship_date for l in plan.lines})
    span = (f"{dates[0]:%m%d}-{dates[-1]:%m%d}" if len(dates) > 1
            else f"{dates[0]:%m%d}" if dates else '확인필요')
    output_file = generate_output_filename('DN추가', plan.period, span, output_dir)

    def _add_table(writer, sheet_name: str, display_name: str) -> None:
        ws = writer.sheets[sheet_name]
        if ws.max_row < 2:
            return
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        tbl = Table(displayName=display_name, ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tbl)

    preview = build_preview(plan)
    reviews = build_review_table(plan)
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        preview.to_excel(writer, sheet_name='추가분', index=False)
        _add_table(writer, '추가분', '추가분')
        ws = writer.sheets['추가분']
        for col_letter, width in (('A', 16), ('B', 16), ('C', 10), ('D', 40),
                                  ('E', 8), ('F', 10), ('G', 13), ('H', 16),
                                  ('I', 24), ('J', 24)):
            ws.column_dimensions[col_letter].width = width
        for r in range(2, ws.max_row + 1):
            for c in (7, 8):
                ws.cell(r, c).number_format = 'yyyy-mm-dd'

        if len(reviews):
            reviews.to_excel(writer, sheet_name='확인필요', index=False)
            _add_table(writer, '확인필요', '확인필요')
            wsr = writer.sheets['확인필요']
            for col_letter, width in (('A', 12), ('B', 16), ('C', 13),
                                      ('D', 22), ('E', 22), ('F', 60)):
                wsr.column_dimensions[col_letter].width = width
            for r in range(2, wsr.max_row + 1):
                wsr.cell(r, 3).number_format = 'yyyy-mm-dd'

    return output_file


def print_summary(plan: Plan) -> None:
    """콘솔 요약 — 무엇이 들어가고 무엇이 빠지는지"""
    print()
    print(f"DN 출고기록 — {plan.period}")
    print("=" * 62)
    print(f"  출고리스트: {plan.delivery_file.name}")
    dn_ids = plan.dn_ids
    if dn_ids:
        span = dn_ids[0] if len(dn_ids) == 1 else f"{dn_ids[0]} ~ {dn_ids[-1]}"
        print(f"  추가 대상:  DN {len(dn_ids)}건 / {len(plan.lines)}라인  ({span})")
    else:
        print("  추가 대상:  없음")

    if plan.skipped:
        print(f"\n  건너뜀 {len(plan.skipped)}건 (이미 입력됨)")
        for r in plan.skipped[:10]:
            date = f"{r.ship_date:%Y-%m-%d}" if r.ship_date is not None else '-'
            print(f"    {r.so_id}  {date}  {r.customer}")
        if len(plan.skipped) > 10:
            print(f"    ... 외 {len(plan.skipped) - 10}건")

    if plan.reviews:
        print(f"\n  {MSG_NOTICE} 확인 필요 {len(plan.reviews)}건 — 자동 입력에서 제외")
        for r in plan.reviews:
            date = f"{r.ship_date:%Y-%m-%d}" if r.ship_date is not None else '-'
            print(f"    {r.so_id}  {date}  [{r.reason}]")
            print(f"        {r.detail}")

    if dn_ids:
        print()
        print(f"  {'DN_ID':<16}{'SO_ID':<16}{'출고일':<12}{'라인':>4}{'수량':>7}  Customer")
        print("  " + "-" * 70)
        for dn_id in dn_ids:
            lines = [l for l in plan.lines if l.dn_id == dn_id]
            head = lines[0]
            print(f"  {dn_id:<16}{head.so_id:<16}"
                  f"{head.ship_date:%Y-%m-%d}  {len(lines):>4}"
                  f"{sum(l.qty for l in lines):>7}  {head.customer_name}")


def create_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog='create_dn',
        description='DN 출고기록 자동 입력 — 공장 출고리스트를 DN_국내에 기록',
        epilog='예시: python create_dn.py P08',
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument('period', help='기간 코드 (예: P08)')
    parser.add_argument(
        '--dry-run', action='store_true',
        help='미리보기 파일만 만들고 워크북은 건드리지 않음')
    parser.add_argument(
        '--yes', '-y', action='store_true',
        help='확인 없이 워크북에 추가 (배치용)')
    parser.add_argument(
        '--no-tax-date', action='store_true',
        help='세금계산서 발행일을 채우지 않음 (기본: 출고일과 동일)')
    parser.add_argument('-v', '--verbose', action='store_true', help='상세 로그')
    return parser


def main() -> int:
    parser = create_argument_parser()
    args = parser.parse_args()
    setup_logging(verbose=args.verbose)
    period = args.period.strip().upper()

    if not NOAH_SO_PO_DN_FILE.exists():
        print(f"{MSG_ERROR} NOAH_SO_PO_DN.xlsx를 찾을 수 없습니다: {NOAH_SO_PO_DN_FILE}")
        return 1

    delivery_file = find_delivery_file(RECON_DIR, period)
    if delivery_file is None:
        print(f"{MSG_ERROR} 출고리스트를 찾을 수 없습니다 "
              f"(po_reconciliation/**/{period}/*리스트*.xlsx)")
        return 1

    try:
        delivery = load_delivery(delivery_file)
        # 반드시 핸들을 닫고 넘어간다 — 열어 두면 뒤에서 Excel이 같은 파일을
        # 쓰기로 못 연다 (우리 프로세스가 우리를 막는다)
        so_df, po_df, dn_df = load_source_frames(
            NOAH_SO_PO_DN_FILE,
            (SO_DOMESTIC_SHEET, PO_DOMESTIC_SHEET, DN_DOMESTIC_SHEET))
    except Exception as e:
        print(f"{MSG_ERROR} 데이터 로드 실패: {e}")
        return 1

    try:
        plan = build_plan(
            period, delivery, so_df, po_df, dn_df,
            delivery_file=delivery_file,
            fill_tax_date=not args.no_tax_date,
        )
    except ValueError as e:
        print(f"{MSG_ERROR} {e}")
        return 1

    print_summary(plan)

    if not plan.lines:
        print("\n추가할 행이 없습니다.")
        if plan.reviews:
            preview_file = write_preview(plan, DN_OUTPUT_DIR)
            print(f"확인 필요 목록: {preview_file}")
        return 0

    preview_file = write_preview(plan, DN_OUTPUT_DIR)
    print(f"\n미리보기: {preview_file}")

    if args.dry_run:
        print("--dry-run — 워크북은 건드리지 않았습니다.")
        return 0

    if not args.yes:
        if not sys.stdin.isatty():
            print(f"\n{MSG_NOTICE} 비대화형 실행이라 워크북에 쓰지 않았습니다 "
                  "(--yes를 주면 바로 씁니다).")
            return 0
        print(f"\n{DN_DOMESTIC_SHEET}에 {len(plan.lines)}행을 추가합니다.")
        if not confirm("진행할까요? (y/N): "):
            print("취소했습니다.")
            return 0

    # xlwings는 여기서만 import — 미리보기·dry-run 경로에 COM 비용을 물리지 않는다
    from po_generator.dn_writer import append_lines

    try:
        result = append_lines(
            NOAH_SO_PO_DN_FILE, DN_DOMESTIC_SHEET, plan.lines,
            backup_dir=DN_BACKUP_DIR,
        )
    except Exception as e:
        print(f"{MSG_ERROR} 워크북 쓰기 실패: {e}")
        return 1

    print(f"\n완료: {result.rows_added}행 추가 "
          f"({DN_DOMESTIC_SHEET} {result.first_row}~{result.last_row}행, "
          f"표 범위 {result.table_ref})")
    if result.backup:
        print(f"백업: {result.backup}")
    print(f"\n{MSG_NOTICE} Excel에서 파워쿼리를 새로고침하면 SO_통합·Order_book에 반영됩니다.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
