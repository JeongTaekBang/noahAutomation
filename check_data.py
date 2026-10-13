#!/usr/bin/env python
"""
NOAH 데이터 정합성 검사
=======================

DB(noah_data.db)를 읽어 **0이어야 하는 것**이 정말 0인지 본다 — 과다출고, 단가 불일치,
SO에 없는 DN 라인, 취소·보류 SO 라인의 DN, PO에 없는 SO, SO에 없는 PO, 한 DN에 거래처 둘,
환율 없는 외화 DN.
검사 목록과 산식은 `po_generator/integrity.py`가 소유한다 (대시보드 'SO/DN 불일치'
카드와 같은 함수). 처리할 일 목록(세금계산서 미발행 등)은 대시보드에 있다.

DB만 읽으므로 대시보드와 숫자가 같다. 대신 DB가 Excel보다 낡았으면 결과도 낡는다 —
그래서 마지막 동기화 시각을 먼저 보여주고, Excel이 그 뒤에 수정됐으면 알린다.
고치는 건 사람 몫이다 — 이 도구는 시트를 건드리지 않는다.

사용법:
    python check_data.py        # 콘솔 요약. 문제가 있으면 generated_check/ 에 상세 xlsx
    python check_data.py -v     # 상세 로그

종료 코드: 0 = 모두 정상, 1 = 문제가 있거나 검사를 못 함
"""

from __future__ import annotations

import argparse
import numbers
import re
import sqlite3
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

import pandas as pd

from po_generator.config import CHECK_OUTPUT_DIR, DB_FILE, NOAH_SO_PO_DN_FILE
from po_generator.db_schema import get_sync_metadata
from po_generator.integrity import CheckResult, run_checks
from po_generator.logging_config import setup_logging

# 콘솔에는 검사당 이만큼만 — 나머지는 상세 파일에 있다
CONSOLE_ROW_LIMIT = 10


def _disp_width(s: str) -> int:
    """문자열의 터미널 표시폭 — 전각(W/F) 문자는 2칸으로 계산."""
    return sum(2 if unicodedata.east_asian_width(ch) in ('W', 'F') else 1 for ch in str(s))


# === DB 읽기 ================================================================

def open_snapshot(db_file: Path) -> sqlite3.Connection:
    """원본 DB의 한 시점 사본을 메모리에 떠서 연다.

    DB는 OneDrive 폴더의 WAL DB라, 원본을 붙잡고 있으면 OneDrive가 파일을 교체하지
    못해 '충돌된 사본'이 생길 수 있다 (`dashboard._db_snapshot`과 같은 이유).
    원본에 거는 잠금은 복사하는 순간뿐이다. 읽기 전용(mode=ro)으로 열지 않는 것도
    대시보드와 같다 — WAL DB는 -shm이 필요한데 읽기 전용 연결은 그걸 만들지 못한다.
    backup()은 읽기만 한다.
    """
    src = sqlite3.connect(str(db_file))
    try:
        snap = sqlite3.connect(':memory:')
        src.backup(snap)
    finally:
        src.close()
    return snap


def last_sync_time(conn: sqlite3.Connection) -> datetime | None:
    """마지막 동기화 시각 — 대시보드 사이드바 '데이터 기준'과 같은 값 (테이블별 최신 중 최댓값)"""
    stamps = []
    for meta in get_sync_metadata(conn).values():
        try:
            stamps.append(datetime.fromisoformat(meta['last_sync']))
        except (TypeError, ValueError):
            continue
    return max(stamps, default=None)


def print_freshness(synced: datetime | None) -> None:
    """데이터 기준 시각 + Excel이 그 뒤에 수정됐으면 경고

    DB가 낡으면 이미 고친 문제가 계속 보이거나, 방금 생긴 문제가 안 보인다.
    """
    if synced is None:
        print("데이터 기준: (동기화 기록 없음)")
        return
    print(f"데이터 기준: {synced:%Y-%m-%d %H:%M} (마지막 DB 동기화)")
    try:
        edited = datetime.fromtimestamp(NOAH_SO_PO_DN_FILE.stat().st_mtime)
    except OSError:
        return
    if edited > synced:
        print(f"[주의] Excel이 마지막 동기화 이후 수정됐습니다 ({edited:%Y-%m-%d %H:%M}).")
        print("       그 뒤에 고친 내용은 아직 DB에 없습니다 — 먼저 DB Sync를 돌리세요 "
              "(메뉴 [8] / python sync_db.py).")


# === 콘솔 출력 ==============================================================

def _status(result: CheckResult) -> str:
    if result.error:
        return "실패"
    return "정상" if result.ok else "문제"


def _cell(value) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ''
    if isinstance(value, numbers.Number) and not isinstance(value, bool):
        return f"{value:,.0f}" if float(value).is_integer() else f"{value:,.2f}"
    return str(value)


def format_table(df: pd.DataFrame) -> list[str]:
    """한글 폭을 맞춘 표 — 글자는 왼쪽, 숫자는 오른쪽 정렬

    pandas `to_string`은 머리글까지 오른쪽으로 붙여 넓은 글자 열에서 머리글이 값과
    떨어져 보인다.
    """
    cells = [[_cell(v) for v in row] for row in df.itertuples(index=False)]
    numeric = [pd.api.types.is_numeric_dtype(df[c]) for c in df.columns]
    widths = [max([_disp_width(c)] + [_disp_width(r[i]) for r in cells])
              for i, c in enumerate(df.columns)]

    def line(values: list[str]) -> str:
        out = []
        for v, w, num in zip(values, widths, numeric):
            pad = ' ' * (w - _disp_width(v))
            out.append(pad + v if num else v + pad)
        return '  '.join(out).rstrip()

    return [line([str(c) for c in df.columns]), *(line(r) for r in cells)]


def print_results(results: list[CheckResult]) -> None:
    width = max(_disp_width(r.check.title) for r in results)
    for r in results:
        title = r.check.title + ' ' * (width - _disp_width(r.check.title))
        count = '-' if r.error else f"{r.count}건"
        print(f"  [{_status(r)}] {title}  {count}")

    for r in results:
        if r.ok:
            continue
        print()
        print('-' * 60)
        head = f"[{_status(r)}] {r.check.title}"
        print(head if r.error else f"{head} ({r.count}건)")
        # 문장 단위로 끊는다 — 글자 수로 감으면 한글은 폭이 두 배라 들쭉날쭉하다
        for sentence in re.split(r'(?<=\.)\s+', r.check.why):
            print(f"  {sentence}")
        if r.error:
            print(f"  오류: {r.error}")
            continue
        print()
        for line in format_table(r.rows.head(CONSOLE_ROW_LIMIT)):
            print(f"  {line}")
        if r.count > CONSOLE_ROW_LIMIT:
            print(f"  ... 외 {r.count - CONSOLE_ROW_LIMIT}건은 상세 파일에")


# === 상세 파일 ==============================================================

def _sheet_name(title: str) -> str:
    """시트 이름에 못 쓰는 문자(`[]:*?/\\`)를 바꾸고 31자로 자른다"""
    for ch in '[]:*?/\\':
        title = title.replace(ch, '-')
    return title[:31]


def write_report(results: list[CheckResult], synced: datetime | None, path: Path) -> None:
    """요약 시트 + 걸린 검사마다 한 장. 문제가 있을 때만 부른다."""
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    summary = pd.DataFrame([{
        '검사': r.check.title,
        '상태': _status(r),
        '건수': None if r.error else r.count,
        '설명': r.check.why,
        '오류': r.error or '',
    } for r in results])

    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        sheets = [('요약', summary)]
        sheets += [(_sheet_name(r.check.title), r.rows)
                   for r in results if r.rows is not None and not r.rows.empty]
        for name, df in sheets:
            df.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = 'A2'
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for idx, col in enumerate(df.columns, start=1):
                values = [col, *df[col].astype(str).head(500)]
                width = min(max(_disp_width(v) for v in values) + 2, 80)
                ws.column_dimensions[get_column_letter(idx)].width = width

        # 설명은 문장이라 폭 제한(80)에 걸린다 — 잘리지 않게 줄바꿈
        head = writer.sheets['요약']
        for row in head.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        head.oddHeader.left.text = "NOAH 데이터 정합성 검사"
        if synced is not None:
            head.oddHeader.right.text = f"데이터 기준 {synced:%Y-%m-%d %H:%M}"


# === CLI ====================================================================

def create_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog='check_data',
        description='NOAH 데이터 정합성 검사 — 0이어야 하는 것이 0인지 DB에서 확인',
        epilog='예시: python check_data.py',
    )
    parser.add_argument('-v', '--verbose', action='store_true', help='상세 로그 출력')
    return parser


def main() -> int:
    args = create_argument_parser().parse_args()
    setup_logging(verbose=args.verbose)

    print("\nNOAH 데이터 정합성 검사")
    print("=" * 60)

    # sqlite3.connect는 없는 파일을 새로 만든다 — 빈 DB로 '모두 정상'을 내면 안 된다
    if not DB_FILE.exists():
        print(f"[오류] DB 파일이 없습니다: {DB_FILE}")
        print("       먼저 DB Sync를 실행하세요 (메뉴 [8] / python sync_db.py).")
        return 1

    print(f"DB: {DB_FILE}")
    try:
        conn = open_snapshot(DB_FILE)
    except sqlite3.Error as e:
        print(f"[오류] DB를 열 수 없습니다: {e}")
        return 1
    try:
        synced = last_sync_time(conn)
        print_freshness(synced)
        results = run_checks(conn)
    finally:
        conn.close()

    print()
    print_results(results)
    print()
    print("=" * 60)

    bad = [r for r in results if not r.ok]
    if not bad:
        print(f"모두 정상 (검사 {len(results)}개)")
        return 0

    failed = sum(1 for r in bad if r.error)
    parts = [f"문제 {len(bad) - failed}개"] + ([f"실패 {failed}개"] if failed else [])
    print(f"{' / '.join(parts)} (검사 {len(results)}개 중)")

    path = CHECK_OUTPUT_DIR / f"정합성검사_{datetime.now():%y%m%d_%H%M%S}.xlsx"
    try:
        write_report(results, synced, path)
    except OSError as e:
        print(f"[경고] 상세 파일 저장 실패: {e}")
    else:
        print(f"상세: {path}")
    return 1


if __name__ == "__main__":
    sys.exit(main())
