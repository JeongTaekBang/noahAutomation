"""원산지증명서(COO) — 값 정하기 · 템플릿 모양 · 서비스 규칙 (Excel COM 없이)

- `build_coo_lines` / `build_coo_header`: 생성기의 순수 함수부 (Excel에 쓰기 전 값)
- 템플릿: **양식은 템플릿 파일이 소유한다** — 생성기의 셀 매핑이 실물과 맞는지 openpyxl로 본다
- 서비스: HS 규칙(일반 = 수출신고 코드 / SECTORIEL = 고객 코드 한 장), PDF 실패 처리.
  조회는 진짜 `FinderService`, 데이터 로드·생성·PDF 변환만 대역 (test_fi_filename과 같은 방식)
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest

from po_generator import coo_generator as coo
from po_generator.config import DEFAULT_HS_CODE
from po_generator.hs_code import HS_COLUMN, HS_ELECTRIC_ACTUATOR, OS_NAME_COLUMN
from po_generator.services import document_service
from po_generator.services.document_service import DocumentService
from po_generator.services.finder_service import FinderService

TEMPLATE = Path(__file__).resolve().parent.parent / 'templates' / 'certificate_of_origin.xlsx'


def item(so='SOO-2026-0001', line=1, name='NA015 230VAC ALS 17SQ', qty=5,
         model='023150', hs='8481.90.0000', po='PO-1', method='OCEAN') -> dict:
    """`_enrich_from_so_lines` + `enrich_hs_codes`를 거친 DN 아이템 한 행"""
    return {
        'SO_ID': so, 'Line item': line, 'Item': name, 'Qty': qty,
        'Model number': model, HS_COLUMN: hs, 'Customer PO': po, 'Shipping Method': method,
    }


def lines_of(rows: list[dict]) -> list[coo.CooLine]:
    return coo.build_coo_lines(pd.DataFrame(rows))


class TestBuildLines:
    def test_운임_줄은_싣지_않는다(self):
        """원산지를 증명할 물건이 아니다 — CI는 HS 빈칸으로 싣지만 COO는 뺀다"""
        lines = lines_of([item(line=1), item(line=2, name='FEDEX COST', hs='')])
        assert [line.line for line in lines] == ['1']

    def test_Customer_PO_다음_SO_다음_SO_라인_순이고_번호는_순번(self):
        """업체는 PO별로 묶어 읽는다. 번호는 정렬이 끝난 뒤 1부터"""
        lines = lines_of([
            item(so='SOO-2026-0298', line=1, po='PO-B'),
            item(so='SOO-2026-0297', line=10, po='PO-A'),
            item(so='SOO-2026-0297', line=2, po='PO-A'),
            item(so='SOO-2026-0298', line=2, po='PO-B'),
        ])
        assert [(line.no, line.customer_po, line.line) for line in lines] == [
            (1, 'PO-A', '2'), (2, 'PO-A', '10'), (3, 'PO-B', '1'), (4, 'PO-B', '2'),
        ]

    def test_한_PO를_SO_여럿으로_받아도_번호가_겹치지_않는다(self):
        """실측: Watergates PO WGDBS2601233을 SO 5개로 받았다 — SO 라인 번호를 찍으면
        DNO-2026-0020에 'PO / Line 1'이 다섯 번 나온다. 순번은 겹치지 않는다"""
        lines = lines_of([
            item(so='SOO-2026-0013', line=1, po='WGDBS2601233'),
            item(so='SOO-2026-0009', line=1, po='WGDBS2601233'),
            item(so='SOO-2026-0009', line=2, po='WGDBS2601233'),
        ])
        assert [(line.no, line.order_no, line.line) for line in lines] == [
            (1, 'SOO-2026-0009', '1'), (2, 'SOO-2026-0009', '2'), (3, 'SOO-2026-0013', '1'),
        ]

    def test_SO_라인은_정수_텍스트(self):
        """정렬에 쓰는 SO Line item — pandas가 float로 읽어도 '3.0'이 아니라 3번째로 줄 선다"""
        assert lines_of([item(line=3.0)])[0].line == '3'

    def test_Item_No는_앞자리_0을_지킨다(self):
        """Model number는 고객 품번이다 — '023150'이 23150이 되면 고객이 못 찾는다"""
        assert lines_of([item(model='023150')])[0].item_no == '023150'
        assert lines_of([item(model=1000057706.0)])[0].item_no == '1000057706'

    def test_HS는_행에_붙은_열에서_온다(self):
        lines = lines_of([item(hs='85013100'), item(line=2, hs=None)])
        assert [line.hs_code for line in lines] == ['85013100', '']

    def test_수량은_정수로(self):
        assert lines_of([item(qty=5.0)])[0].qty == 5

    @pytest.mark.parametrize('qty', [None, 'abc'])
    def test_수량을_못_읽으면_멈춘다(self, qty):
        """증명서에 추정 수량을 찍지 않는다 (CI는 1로 메운다)"""
        with pytest.raises(ValueError, match='수량'):
            lines_of([item(qty=qty)])

    def test_cells는_머리글_순서(self):
        line = lines_of([item()])[0]
        assert line.cells() == ['1', 'PO-1', '023150', 'NA015 230VAC ALS 17SQ',
                                '8481.90.0000', 5]
        assert len(line.cells()) == len(coo.ITEM_HEADERS)


class TestBuildHeader:
    BASE = {
        'DN_ID': 'DNO-2026-0190', 'Customer name': 'WATERGATES GMBH',
        'Bill to 1': 'Watergates GmbH & Co.KG',
        'Bill to 2': 'Obebecksener Str. 78 32547 Bad Oeynhausen',
        'Bill to 3': 'GERMANY', 'Customer country': 'GERMANY',
        '출고일': pd.Timestamp('2026-10-02'),
    }

    def header(self, lines=None, **overrides) -> coo.CooHeader:
        data = pd.Series({**self.BASE, **overrides})
        return coo.build_coo_header(data, lines if lines is not None else lines_of([item()]))

    def test_증명서_번호와_인보이스_번호는_DN에서(self):
        h = self.header()
        assert h.certificate_no == 'COO-DNO-2026-0190'
        assert h.invoice_no == 'DNO-2026-0190'

    def test_발행일은_CI_Invoice_Date와_같은_출고일(self):
        assert self.header().issue_date == '2026-10-02'

    def test_출고일이_없으면_오늘(self):
        assert self.header(**{'출고일': None}).issue_date == datetime.now().strftime('%Y-%m-%d')

    def test_수하인은_Bill_to_빈_줄은_건너뛴다(self):
        h = self.header(**{'Bill to 2': None})
        assert h.consignee == ('Watergates GmbH & Co.KG', 'GERMANY')

    def test_Bill_to가_다_비면_고객명이라도(self):
        """수하인이 빈 증명서는 내보내지 않는다"""
        h = self.header(**{'Bill to 1': None, 'Bill to 2': '', 'Bill to 3': None})
        assert h.consignee == ('WATERGATES GMBH',)

    def test_도착국은_고객_마스터_Country_우선(self):
        """Bill to 3은 거래처마다 국가가 아니라 주소 끝줄이기도 하다"""
        assert self.header(**{'Bill to 3': '45 RUE X'}).destination == 'GERMANY'
        assert self.header(**{'Customer country': None, 'Bill to 3': 'France'}).destination == 'France'

    def test_번호_목록은_표에_실린_줄에서_순서대로_중복_없이(self):
        lines = lines_of([
            item(so='SOO-2026-0297', po='WG-1'), item(so='SOO-2026-0297', line=2, po='WG-1'),
            item(so='SOO-2026-0298', po='WG-2'),
        ])
        h = self.header(lines=lines)
        assert h.order_nos == 'SOO-2026-0297, SOO-2026-0298'
        assert h.customer_pos == 'WG-1, WG-2'

    def test_운송은_방법과_선적항(self):
        assert self.header().transport == 'OCEAN, from INCHEON, KOREA'
        two = lines_of([item(method='OCEAN'), item(line=2, method='AIR')])
        assert self.header(lines=two).transport == 'OCEAN / AIR, from INCHEON, KOREA'
        none = lines_of([item(method='')])
        assert self.header(lines=none).transport == 'From INCHEON, KOREA'

    def test_참조_값_순서(self):
        h = self.header()
        assert h.references() == [h.invoice_no, h.order_nos, h.customer_pos,
                                  h.transport, h.destination]


class TestTemplateShape:
    """생성기 셀 매핑 ↔ 템플릿 실물 — 한쪽만 바뀌면 값이 엉뚱한 칸에 조용히 들어간다"""

    @pytest.fixture(scope='class')
    def ws(self):
        from openpyxl import load_workbook
        return load_workbook(TEMPLATE).worksheets[0]

    @staticmethod
    def _merge_of(ws, ref: str) -> str | None:
        for rng in ws.merged_cells.ranges:
            if ref in rng:
                return str(rng)
        return None

    def test_표_머리글(self, ws):
        row = coo.HEADER_ROW
        actual = tuple(ws.cell(row=row, column=c).value for c in range(1, len(coo.ITEM_HEADERS) + 1))
        assert actual == coo.ITEM_HEADERS

    def test_Total_위치(self, ws):
        assert ws[f'{coo.TOTAL_LABEL_COL}{coo.ITEM_START_ROW + 1}'].value == 'Total'

    def test_품목_행에는_병합이_없다(self, ws):
        """병합이 없어야 rows.autofit()이 먹고 행 복사 삽입이 병합을 잃지 않는다
        (coo_generator docstring — CI·PL과 다른 점)"""
        row = coo.ITEM_START_ROW
        merged = [str(r) for r in ws.merged_cells.ranges if r.min_row <= row <= r.max_row]
        assert not merged

    def test_텍스트_칸은_텍스트_서식(self, ws):
        """Order No.·Line·Item No.·HS — '023150'이 숫자로 바뀌면 앞 0이 사라진다"""
        for col in 'ABCE':
            assert ws[f'{col}{coo.ITEM_START_ROW}'].number_format == '@', col

    @pytest.mark.parametrize('label_row,label', [
        (7, 'Certificate No.:'),
        (8, 'Date of Issue:'),
        (15, 'Consignee:'),
        (19, 'Commercial Invoice No.:'),
        (20, 'Order No.:'),
        (21, 'Customer Order No.:'),
        (22, 'Means of Transport:'),
        (23, 'Country of Destination:'),
    ])
    def test_값_칸이_라벨_옆_병합의_좌상단(self, ws, label_row, label):
        assert ws[f'A{label_row}'].value == label
        assert self._merge_of(ws, f'C{label_row}') == f'C{label_row}:F{label_row}'

    def test_셀_매핑_상수가_라벨과_맞다(self):
        assert coo.CELL_CERTIFICATE_NO == 'C7'
        assert coo.CELL_ISSUE_DATE == 'C8'
        assert coo.CONSIGNEE_START_ROW == 15
        assert coo.REFERENCE_START_ROW == 19

    def test_수하인_4줄이_모두_병합_칸(self, ws):
        for offset in range(coo.CONSIGNEE_ROW_COUNT):
            row = coo.CONSIGNEE_START_ROW + offset
            assert self._merge_of(ws, f'C{row}') == f'C{row}:F{row}'

    def test_표_머리글이_페이지마다_반복된다(self, ws):
        assert ws.print_title_rows.replace('$', '') == f'{coo.HEADER_ROW}:{coo.HEADER_ROW}'

    def test_맺음_블록_끝이_템플릿_마지막_행(self, ws):
        """생성기는 이 행까지를 한 장에 모은다 — 템플릿 끝이 바뀌면 상수도 같이.
        서명 이미지는 그 블록 안(For and on behalf 아래 ~ 마지막 행)에 걸려 있어야 같이 넘어간다"""
        assert ws.print_area.replace('$', '').endswith(f'F{coo.TEMPLATE_LAST_ROW}')
        on_behalf = next(r for r in range(1, coo.TEMPLATE_LAST_ROW + 1)
                         if str(ws[f'A{r}'].value or '').startswith('For and on behalf of'))
        sign_rows = [img.anchor._from.row + 1 for img in ws._images]  # 0-based → 1-based
        assert any(on_behalf < r <= coo.TEMPLATE_LAST_ROW for r in sign_rows)

    def test_폭만_한_장에_맞추고_높이는_자동(self, ws):
        """고정 배율은 프린터마다 인쇄 폭이 달라 열이 밀린다(doPDF 100% → Qty 열이 다음 장).
        높이까지 장수로 고정하면 Excel이 수동 나누기를 무시해 서명 블록 유지가 안 먹는다"""
        assert ws.sheet_properties.pageSetUpPr.fitToPage
        assert ws.page_setup.fitToWidth == 1
        assert ws.page_setup.fitToHeight == 0

    def test_열면_인쇄영역이_보인다(self, ws):
        """CI·PL 템플릿과 같은 '페이지 나누기 미리 보기' — 파란 테두리로 인쇄영역이 바로 보인다.
        생성 파일은 템플릿 사본이라 같이 따라간다"""
        assert ws.sheet_view.view == 'pageBreakPreview'

    def test_기본_글꼴이_Arial_10(self, ws):
        """빈 칸의 글꼴이 행 높이 하한을 정한다 — 테마 글꼴이면 한글 Excel에서 맑은 고딕 11
        (16.5pt)이 되어 헤더 행이 전부 18.5pt로 부풀었다 (실측)"""
        normal = ws.parent._named_styles['Normal'].font
        assert (normal.name, normal.sz) == ('Arial', 10)


    def test_선언은_수출자_RCK_생산은_NOAH(self, ws):
        """2026-10-14 결정 — RCK(CI의 Seller)가 선언·서명하고 생산 공장 NOAH를 밝힌다"""
        values = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        assert (ws['A10'].value, ws['C10'].value) == ('Exporter:', 'Rotork Controls Korea Co., Ltd.')
        assert (ws['A13'].value, ws['C13'].value) == ('Manufacturer:', 'NOAH Actuation Co., Ltd.')
        assert any(v.startswith('For and on behalf of Rotork Controls Korea') for v in values)
        statement = next(v for v in values if v.startswith('Rotork Controls Korea Co., Ltd., as the exporter'))
        assert 'manufactured by NOAH Actuation Co., Ltd.' in statement

    def test_선언문은_보관_기간을_약속하지_않는다(self, ws):
        """샘플(Bifold) 문구 그대로 — FTA식 '5년 보관'은 사용자 지시로 뺐다 (2026-10-14)"""
        statement = next(c.value for row in ws.iter_rows() for c in row
                         if isinstance(c.value, str) and 'as the exporter' in c.value)
        assert 'year' not in statement.lower()


class TestClosingBlockSplit:
    """맺음 블록(Total ~ 서명) 페이지 판정 — 품목 마지막 행 36, 블록 끝 50"""

    @pytest.mark.parametrize('breaks,expected', [
        ([], False),                 # 한 장
        ([30], False),               # 품목 중간에서 넘어감 — 표가 이어지는 정상 경우
        ([36], False),               # 마지막 품목부터 새 장 — 품목과 블록이 같이 있다
        ([37], True),                # Total부터 새 장 — 품목 없이 머리글·합계만
        ([45], True),                # 선언문 중간
        ([50], True),                # Title만 넘어감 (실측 DNO-2026-0135)
        ([51], False),               # 블록 뒤
        ([20, 45], True),
    ])
    def test_판정(self, breaks, expected):
        assert coo.closing_block_split(breaks, last_item_row=36, block_end_row=50) is expected


# === 서비스 ===================================================================

SECTORIEL = 'SECTORIEL SA'


def dn_row(line=1, name='NA28 280Nm 24s. 230V50HZ', customer='WATERGATES GMBH') -> dict:
    """`load_dn_export_data()` 한 행"""
    return {
        'DN_ID': 'DNO-2026-0190', 'SO_ID': 'SOO-2026-0298', 'Customer name': customer,
        'Customer PO': 'PO-1', 'Item': name, 'Line item': line, 'Qty': 2,
        'Shipping Method': 'OCEAN', 'Bill to 1': customer, '출고일': pd.Timestamp('2026-10-02'),
    }


def so_rows(os_name: str | None = 'Noah NA') -> pd.DataFrame:
    rows = [{'SO_ID': 'SOO-2026-0298', 'Line item': n, 'Model number': f'M{n}'} for n in (1, 2)]
    df = pd.DataFrame(rows)
    if os_name is not None:
        df[OS_NAME_COLUMN] = os_name
    return df


@pytest.fixture
def make_service(tmp_path, monkeypatch):
    """행 목록으로 DocumentService — 생성기는 받은 items_df를 기록하고 빈 파일만 쓴다"""
    template = tmp_path / 'certificate_of_origin.xlsx'
    template.write_bytes(b'template')
    monkeypatch.setattr(document_service, 'COO_TEMPLATE_FILE', template)
    monkeypatch.setattr(document_service, 'COO_OUTPUT_DIR', tmp_path / 'generated_coo')

    captured: dict = {}

    def fake_create(template_path, output_path, order_data, items_df):
        captured['items_df'] = items_df
        output_path.write_bytes(b'dummy')

    def fake_pdf(xlsx_path):
        pdf = Path(xlsx_path).with_suffix('.pdf')
        pdf.write_bytes(b'%PDF')
        return pdf

    monkeypatch.setattr(document_service, 'create_coo_xlwings', fake_create)
    monkeypatch.setattr(document_service, 'export_pdf', fake_pdf)

    def _make(rows: list[dict], so: pd.DataFrame | None = None) -> DocumentService:
        finder = FinderService()
        monkeypatch.setattr(finder, 'load_dn_export_data', lambda: pd.DataFrame(rows))
        so_df = so if so is not None else so_rows()
        monkeypatch.setattr(finder, 'load_so_export_data', lambda: so_df)
        return DocumentService(finder=finder)

    return _make, captured


class TestService:
    def test_xlsx와_PDF를_둘_다_남긴다(self, make_service):
        make, _ = make_service
        result = make([dn_row()]).generate_coo('DNO-2026-0190')
        assert result.success and not result.errors
        assert re.fullmatch(r'COO_DNO-2026-0190_WATERGATES_GMBH_\d{6}\.xlsx', result.output_file.name)
        assert [p.suffix for p in result.extra_files] == ['.pdf']
        assert result.extra_files[0].stem == result.output_file.stem

    def test_일반_거래처는_수출신고_코드(self, make_service):
        make, captured = make_service
        make([dn_row(), dn_row(line=2, name='FEDEX COST')]).generate_coo('DNO-2026-0190')
        assert list(captured['items_df'][HS_COLUMN]) == [DEFAULT_HS_CODE, '']

    def test_SECTORIEL은_고객_코드_한_장만(self, make_service):
        """COO는 고객 수입통관용 — 같이 제출되는 CI-HS와 코드가 같아야 한다.
        수출신고용 판을 따로 만들지 않는다 (xlsx 1 + PDF 1)"""
        make, captured = make_service
        result = make([dn_row(customer=SECTORIEL)]).generate_coo('DNO-2026-0190')
        assert result.success and not result.errors
        assert list(captured['items_df'][HS_COLUMN]) == [HS_ELECTRIC_ACTUATOR]
        assert [p.suffix for p in result.extra_files] == ['.pdf']

    def test_SECTORIEL인데_OS_name이_없으면_만들지_않는다(self, make_service):
        """전 행 HS 빈칸 증명서를 성공으로 내보내지 않는다"""
        make, captured = make_service
        result = make([dn_row(customer=SECTORIEL)], so=so_rows(os_name=None)).generate_coo(
            'DNO-2026-0190')
        assert not result.success
        assert 'items_df' not in captured

    def test_PDF가_실패하면_xlsx는_남기고_오류로_알린다(self, make_service, monkeypatch):
        make, _ = make_service

        def broken_pdf(xlsx_path):
            raise RuntimeError('Excel 없음')

        monkeypatch.setattr(document_service, 'export_pdf', broken_pdf)
        result = make([dn_row()]).generate_coo('DNO-2026-0190')
        assert result.output_file.exists()
        assert result.extra_files == []
        assert result.errors and 'PDF' in result.errors[0]

    def test_없는_DN(self, make_service):
        make, _ = make_service
        assert not make([dn_row()]).generate_coo('DNO-2026-9999').success
