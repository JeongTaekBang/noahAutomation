"""
문서 레이아웃 계산·감시 테스트 (Excel 없이)
=============================================

해외 문서 5종(OC·FI·PI·CI·PL)의 레이아웃 규칙을 COM 없이 지킨다:

- `address_row_heights()` — 주소 블록 행 높이 (순수 함수)
- 레이아웃 상수 불변식 (병합 열·최소 행 높이·여유)
- 생성기들이 공용 헬퍼(`layout_item_rows`/`layout_address_rows`)를 쓰는지 소스 감시

(예전의 "한 페이지 빈 행 채우기" 테스트는 기능 제거와 함께 사라졌다 —
2026-08-05, 1아이템 문서가 빈 격자를 달고 나가 표를 아이템에서 끝내기로 결정.)
"""

import pytest
from pathlib import Path

from po_generator.excel_helpers import (
    HS_COLUMN_LETTER,
    assert_hs_template,
    ITEM_NAME_MERGED_COLS,
    ITEM_NAME_MERGED_COLS_HS,
    MIN_ITEM_ROW_HEIGHT,
    ROW_HEIGHT_PAD,
    _item_name_cols,
    address_row_heights,
)


class TestAddressRowHeights:
    """주소 블록 행 높이 계산 (순수 함수) — 기준값은 OC 템플릿 실측치

    주소 행 3개는 각 15.95pt. 왼쪽은 행별 한 줄 병합(A:E), 오른쪽은 3행을
    세로로 걸친 병합(G13:I15) 하나다.
    """

    CUR = [15.95, 15.95, 15.95]

    def test_짧은_주소는_현재_높이를_유지한다(self):
        """한 줄 측정치(~13pt)+여유가 15.95를 넘지 않으면 모양이 안 변한다"""
        assert address_row_heights(self.CUR, [13.0, 13.0, 13.0], 0.0, pad=2.0) == self.CUR

    def test_접힌_왼쪽_줄만_자란다(self):
        """SECTORIEL 실측: 74자 bill_to_2가 2줄(26pt)로 접히면 그 행만 커진다"""
        heights = address_row_heights(self.CUR, [13.0, 26.0, 13.0], 0.0, pad=2.0)
        assert heights == [15.95, 28.0, 15.95]

    def test_오른쪽_블록_부족분은_균등_분배된다(self):
        """납품 주소가 3행 합(47.85pt)보다 크면 세 행이 같이 자란다"""
        heights = address_row_heights(self.CUR, [13.0] * 3, 60.0, pad=2.0)
        assert sum(heights) == pytest.approx(62.0)
        assert heights[0] == pytest.approx(heights[1]) == pytest.approx(heights[2])

    def test_오른쪽이_행_합에_들어가면_분배하지_않는다(self):
        """2줄 납품 주소(~28pt)는 3행 합 47.85pt에 이미 들어간다 — SECTORIEL 실측"""
        assert address_row_heights(self.CUR, [13.0] * 3, 28.0, pad=2.0) == self.CUR

    def test_왼쪽_성장으로_이미_충분하면_더_키우지_않는다(self):
        heights = address_row_heights(self.CUR, [13.0, 40.0, 13.0], 50.0, pad=2.0)
        assert heights == [15.95, 42.0, 15.95]  # 합 73.9 >= 52

    def test_두_요구를_동시에_만족한다(self):
        """어느 쪽이 이기든 최종 합은 오른쪽 요구 이상, 각 행은 왼쪽 요구 이상"""
        for right in (0.0, 30.0, 55.0, 90.0):
            heights = address_row_heights(self.CUR, [13.0, 26.0, 13.0], right, pad=2.0)
            assert all(h >= n + 2.0 for h, n in zip(heights, [13.0, 26.0, 13.0]))
            assert all(h >= c for h, c in zip(heights, self.CUR))
            if right > 0:
                assert sum(heights) >= right + 2.0 - 1e-9

    def test_길이가_어긋나면_조용히_진행하지_않는다(self):
        with pytest.raises(ValueError):
            address_row_heights([15.95, 15.95], [13.0] * 3, 0.0)

    def test_빈_입력은_빈_결과(self):
        assert address_row_heights([], [], 60.0) == []


class TestGeneratorsLayoutAddresses:
    """주소를 쓰는 생성기(OC·FI)가 공용 주소 레이아웃을 부르는지 감시

    주소 칸은 병합 셀이라 wrap 없이 길면 병합 경계에서 잘린다(2026-08-05 실측
    74자/81자 클립). 값만 쓰고 layout_address_rows를 빼먹으면 재발한다.
    PI·CI·PL은 주소 셀을 채우지 않아 대상이 아니다.
    """

    @pytest.mark.parametrize('module_name', [
        'po_generator.oc_generator',
        'po_generator.fi_generator',
    ])
    def test_주소를_쓰면_layout_address_rows를_부른다(self, module_name):
        import importlib
        import inspect

        source = inspect.getsource(importlib.import_module(module_name))
        assert 'layout_address_rows(' in source, (
            f"{module_name}: 주소 셀을 채우면서 layout_address_rows를 부르지 않는다 — "
            f"긴 주소가 병합 경계에서 잘린다"
        )


class TestLayoutConstants:
    def test_최소_행_높이는_템플릿_기본값(self):
        assert MIN_ITEM_ROW_HEIGHT == 15.0

    def test_여유는_양수다(self):
        """병합 셀은 안쪽 여백이 미세하게 커서 딱 맞추면 마지막 줄이 잘린다"""
        assert ROW_HEIGHT_PAD > 0

    def test_병합_열은_문서_5종_공통이다(self):
        """CI와 PL은 늘 같이 첨부되므로 두 문서가 같은 값을 써야 한다"""
        assert ITEM_NAME_MERGED_COLS == 'ABCD'

    def test_HS판은_D만_떼어낸_A_C다(self):
        """라인별 HS 판의 품목명은 A:C — 떨어져 나온 D가 곧 HS 열이다"""
        assert ITEM_NAME_MERGED_COLS_HS == 'ABC'
        assert ITEM_NAME_MERGED_COLS.startswith(ITEM_NAME_MERGED_COLS_HS)
        assert ITEM_NAME_MERGED_COLS[len(ITEM_NAME_MERGED_COLS_HS):] == HS_COLUMN_LETTER

    def test_판은_둘뿐이고_헬퍼가_고른다(self):
        """호출부는 bool만 넘긴다 — 열 문자열을 직접 넘길 수 있으면 CI와 PL이 갈린다"""
        assert _item_name_cols(False) == ITEM_NAME_MERGED_COLS
        assert _item_name_cols(True) == ITEM_NAME_MERGED_COLS_HS

    def test_생성기들이_병합_열을_재정의하지_않는다(self):
        """한쪽만 바뀌면 같은 봉투 안의 두 장이 다른 줄간격으로 나간다"""
        import po_generator.ci_generator as ci
        import po_generator.fi_generator as fi
        import po_generator.oc_generator as oc
        import po_generator.pi_generator as pi
        import po_generator.pl_generator as pl

        for mod in (oc, fi, pi, ci, pl):
            assert not hasattr(mod, 'ITEM_NAME_MERGED_COLS'), (
                f"{mod.__name__}이 병합 열을 따로 정의하고 있다 — "
                f"excel_helpers.ITEM_NAME_MERGED_COLS 하나만 쓸 것"
            )


class TestGeneratorsUseSharedHelper:
    """다섯 생성기가 `rows.autofit()`으로 돌아가지 않았는지 감시

    병합 셀에는 그게 먹지 않는다 — 오히려 1줄로 줄여 긴 품목명을 잘라낸다.
    """

    @pytest.mark.parametrize('module_name', [
        'po_generator.oc_generator',
        'po_generator.fi_generator',
        'po_generator.pi_generator',
        'po_generator.ci_generator',
        'po_generator.pl_generator',
    ])
    def test_아이템_영역에_rows_autofit을_쓰지_않는다(self, module_name):
        import importlib
        import inspect

        source = inspect.getsource(importlib.import_module(module_name))
        offending = [
            line.strip() for line in source.splitlines()
            if 'rows.autofit()' in line and not line.strip().startswith('#')
        ]
        assert not offending, f"{module_name}: {offending}"

    @pytest.mark.parametrize('module_name', [
        'po_generator.oc_generator',
        'po_generator.fi_generator',
        'po_generator.pi_generator',
        'po_generator.ci_generator',
        'po_generator.pl_generator',
    ])
    def test_복합_헬퍼로_병합과_행높이를_처리한다(self, module_name):
        """layout_item_rows = 병합 보장 + 행 높이 교정 세트 — 한쪽만 부르는 실수 차단"""
        import importlib
        import inspect

        source = inspect.getsource(importlib.import_module(module_name))
        assert 'layout_item_rows(' in source, f"{module_name}: layout_item_rows 미사용"


class TestHsTemplateShape:
    """**양식은 템플릿 파일이 소유한다** — 그 실물을 본다 (COM 없이 openpyxl로)

    예전엔 생성할 때마다 `apply_hs_layout()`이 임시 사본을 고쳤다. 판이 둘(표준/고객용)일
    때는 템플릿을 복제하면 갈라지니 그럴 만했는데, 양식이 하나로 통일되면서 근거가
    사라져 템플릿에 구웠다 (2026-09-07). 그래서 이제 **템플릿이 어긋나면 산출물이 어긋난다**
    — 여기서 직접 감시한다. 구 양식은 `templates/Old/`에 있어 기준값을 하드코딩하지 않아도 된다.
    """

    TEMPLATES = Path(__file__).resolve().parent.parent / 'templates'
    NAMES = ('commercial_invoice.xlsx', 'packing_list.xlsx')
    HEADER_ROW, FIRST_ITEM_ROW = 18, 20

    # 열 하나가 줄면 그 열 몫의 안쪽 여백이 사라져, A:C는 A:D보다 **저장된 폭 숫자가**
    # 패딩 한 몫만큼 작다 (실측 0.83자 — `excel_helpers` 주석). 화면·인쇄 폭은 같다.
    COLUMN_PADDING = 0.83

    @staticmethod
    def _sheet(path):
        from openpyxl import load_workbook

        return load_workbook(path).worksheets[0]

    @classmethod
    def _total_row(cls, ws):
        """아이템 격자의 끝 — 그 아래 Shipping Mark 등은 병합 규칙이 다르다"""
        for r in range(cls.FIRST_ITEM_ROW, cls.FIRST_ITEM_ROW + 30):
            if str(ws.cell(r, 1).value or '').strip() == 'Total':
                return r
        raise AssertionError('Total 행을 찾지 못했다')

    @pytest.fixture(params=NAMES)
    def name(self, request):
        if not (self.TEMPLATES / request.param).exists():
            pytest.skip(f"템플릿 없음: {request.param}")
        return request.param

    def test_HS_열_머리글이_있다(self, name):
        """생성기가 이걸 보고 구 양식 혼입을 막는다 (`assert_hs_template`)"""
        ws = self._sheet(self.TEMPLATES / name)
        assert ws[f'{HS_COLUMN_LETTER}{self.HEADER_ROW}'].value == 'HS CODE'

    def test_품목명은_HS_열을_비켜_병합돼_있다(self, name):
        """A:D로 남아 있으면 D에 쓴 코드가 병합 셀에 삼켜져 조용히 사라진다"""
        ws = self._sheet(self.TEMPLATES / name)
        grid = range(self.HEADER_ROW, self._total_row(ws) + 1)
        merges = {
            str(m) for m in ws.merged_cells.ranges
            if m.min_col == 1 and m.min_row in grid
        }
        assert len(merges) == len(grid), f"{name}: 격자 행마다 병합이 있어야 한다 — {sorted(merges)}"
        last = ITEM_NAME_MERGED_COLS_HS[-1]
        assert all(m.split(':')[1].startswith(last) for m in merges), sorted(merges)[:5]

    def test_문서_단위_HS는_비어_있다(self, name):
        """라인별 코드와 한 장에 같이 남으면 문서가 두 HS를 주장한다"""
        ws = self._sheet(self.TEMPLATES / name)
        assert ws['H12'].value in (None, ''), ws['H12'].value
        assert ws['I12'].value in (None, ''), ws['I12'].value

    def test_인쇄는_한_페이지_폭에_맞춘다(self, name):
        """D를 넣느라 늘어난 폭(약 12자)을 배율로 흡수한다 — 열 폭을 안 건드리려고"""
        ws = self._sheet(self.TEMPLATES / name)
        assert ws.sheet_properties.pageSetUpPr.fitToPage is True
        assert ws.page_setup.fitToWidth in (1, '1', None), ws.page_setup.fitToWidth

    # 구 양식의 품목명 폭(A:D) 실측 — D를 HS 열로 떼어내기 전 값.
    # `templates/Old/`에 보관본이 있지만 그 폴더는 .gitignore 대상이라 다른 환경에는
    # 없다. 기준이 사라지면 감시도 사라지므로 여기 적어 둔다.
    OLD_ITEM_NAME_WIDTH = {'commercial_invoice.xlsx': 37.33, 'packing_list.xlsx': 36.66}

    def test_품목명_폭이_구_양식보다_좁지_않다(self, name):
        """좁히면 200줄이 406 → 505줄로 늘어 SECTORIEL 두 장의 쪽수가 갈린다 (실측)"""
        ws = self._sheet(self.TEMPLATES / name)
        now = sum(
            getattr(ws.column_dimensions.get(c), 'width', 0) or 0
            for c in ITEM_NAME_MERGED_COLS_HS
        )
        was = self.OLD_ITEM_NAME_WIDTH[name]
        assert now >= was - self.COLUMN_PADDING, f"{name}: 품목명 폭 {was} -> {now}"

    def test_구_양식_보관본이_있으면_구_양식_그대로다(self, name):
        """`templates/Old/`가 새 양식으로 덮이면 되돌릴 곳이 없어진다 (있을 때만 확인)"""
        old_path = self.TEMPLATES / 'Old' / name
        if not old_path.exists():
            pytest.skip(f"구 양식 보관본 없음 (.gitignore 대상): Old/{name}")
        ws = self._sheet(old_path)
        assert ws[f'{HS_COLUMN_LETTER}{self.HEADER_ROW}'].value != 'HS CODE'
        assert ws['I12'].value, "구 양식은 문서 단위 HS를 갖고 있었다"
        was = sum(
            getattr(ws.column_dimensions.get(c), 'width', 0) or 0
            for c in ITEM_NAME_MERGED_COLS
        )
        assert was == pytest.approx(self.OLD_ITEM_NAME_WIDTH[name], abs=0.01), (
            f"{name}: 위 테스트의 기준값이 보관본과 다르다 — {was}"
        )


class TestAssertHsTemplate:
    """구 양식 혼입 가드 — 값 하나만 읽으므로 스텁으로 충분하다"""

    class _Sheet:
        def __init__(self, value):
            self._value = value

        def range(self, _addr):
            return type('R', (), {'value': self._value})()

    def test_HS_머리글이_있으면_통과(self):
        assert_hs_template(self._Sheet('HS CODE'), 18) is None

    @pytest.mark.parametrize('value', [None, '', 'Description', '품목'])
    def test_없으면_멈춘다(self, value):
        with pytest.raises(ValueError, match='HS CODE'):
            assert_hs_template(self._Sheet(value), 18)

    def test_대소문자와_공백은_허용(self):
        assert_hs_template(self._Sheet(' hs code '), 18) is None


class TestGeneratorsDoNotHardcodeColumns:
    """생성기가 병합 열 문자열을 직접 들고 있으면 CI와 PL이 갈린다"""

    @pytest.mark.parametrize('module_name', [
        'po_generator.ci_generator',
        'po_generator.pl_generator',
    ])
    def test_병합_열_리터럴이_없다(self, module_name):
        import importlib
        import inspect

        source = inspect.getsource(importlib.import_module(module_name))
        code_lines = [
            line for line in source.splitlines()
            if not line.strip().startswith('#')
        ]
        for literal in ("'ABCD'", '"ABCD"', "'ABC'", '"ABC"'):
            offending = [ln.strip() for ln in code_lines if literal in ln]
            assert not offending, f"{module_name}: {literal} 하드코딩 — {offending}"
