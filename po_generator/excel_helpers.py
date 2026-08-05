"""
Excel 헬퍼 함수 모듈
====================

여러 모듈에서 공통으로 사용하는 Excel 관련 유틸리티 함수를 제공합니다.

- XlConstants: Excel COM 매직 넘버 상수
- xlwings_app_context: xlwings App 생명주기 관리 컨텍스트 매니저
- prepare_template: 템플릿 파일을 임시 폴더로 복사
- find_item_start_row_openpyxl: openpyxl 워크시트용
- find_item_start_row_xlwings: xlwings 워크시트용
- batch_write_rows: 2D 리스트를 한 번에 쓰기 (성능 최적화)
- batch_read_column: 열의 값을 한 번에 읽기 (성능 최적화)
- delete_rows_range / insert_copied_rows: 연속 행을 한 번에 삭제/삽입 (성능 최적화)
- layout_item_rows: 아이템 행 마무리 손질 = 병합 보장 + 행 높이 교정 (문서 5종 공통 진입점)
- layout_address_rows: 헤더 주소 블록 줄바꿈 + 행 높이 확보 (OC·FI 공통)
"""

from __future__ import annotations

import logging
import shutil
import tempfile
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING, Generator

import xlwings as xw

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet as OpenpyxlWorksheet

from po_generator.config import ITEM_START_ROW_FALLBACK

logger = logging.getLogger(__name__)


# === Excel COM 상수 클래스 ===

class XlConstants:
    """Excel COM 인터페이스 매직 넘버 상수

    xlwings에서 Excel COM API를 직접 호출할 때 사용하는 상수들입니다.
    매직 넘버 대신 이 상수들을 사용하여 코드 가독성을 높입니다.

    참고: https://docs.microsoft.com/en-us/office/vba/api/excel.xlshiftdirection
    """
    # Shift 방향
    xlShiftUp = -4162      # 행 삭제 시 아래 행이 위로 올라옴
    xlShiftDown = -4121    # 행 삽입 시 기존 행이 아래로 내려옴

    # 테두리 위치
    xlEdgeLeft = 7         # 왼쪽 테두리
    xlEdgeTop = 8          # 상단 테두리
    xlEdgeBottom = 9       # 하단 테두리
    xlEdgeRight = 10       # 오른쪽 테두리
    xlInsideVertical = 11  # 내부 세로선
    xlInsideHorizontal = 12  # 내부 가로선

    # 테두리 스타일
    xlContinuous = 1       # 실선
    xlNone = -4142         # 테두리 없음
    xlThin = 2             # 얇은 선
    xlMedium = -4138       # 중간 두께

    # 정렬 (XlVAlign / XlHAlign 공용)
    xlCenter = -4108       # 가운데 정렬


# === xlwings 앱 컨텍스트 매니저 ===

@contextmanager
def xlwings_app_context(
    visible: bool = False,
    display_alerts: bool = False,
    screen_updating: bool = False,
) -> Generator[xw.App, None, None]:
    """xlwings App 생명주기를 안전하게 관리하는 컨텍스트 매니저

    리소스 누수 방지를 위해 오류 발생 시에도 Excel 프로세스를 정리합니다.

    Args:
        visible: Excel 창 표시 여부 (기본: False)
        display_alerts: Excel 알림 표시 여부 (기본: False)
        screen_updating: 화면 업데이트 여부 (기본: False, 성능 향상)

    Yields:
        xw.App: xlwings App 객체

    Example:
        with xlwings_app_context() as app:
            wb = app.books.open(str(template_path))
            ws = wb.sheets[0]
            # ... 작업 수행 ...
            wb.save(str(output_path))
        # 컨텍스트 종료 시 자동으로 정리됨
    """
    app = None
    try:
        app = xw.App(visible=visible)
        app.display_alerts = display_alerts
        app.screen_updating = screen_updating
        yield app
    finally:
        if app is not None:
            # 모든 워크북 닫기 시도
            try:
                for wb in app.books:
                    try:
                        wb.close()
                    except Exception:
                        pass
            except Exception:
                pass
            # App 종료
            try:
                app.quit()
            except Exception:
                pass


# === 템플릿 준비 헬퍼 ===

def prepare_template(template_path: Path, prefix: str = "template") -> tuple[Path, Path]:
    """템플릿 파일을 임시 디렉토리에 복사하고 경로 반환

    xlwings COM 인터페이스는 한글 경로에서 문제가 발생할 수 있습니다.
    템플릿을 임시 폴더로 복사하여 이 문제를 우회합니다.

    Args:
        template_path: 원본 템플릿 파일 경로
        prefix: 임시 파일 접두사 (기본: "template")

    Returns:
        (temp_template_path, temp_output_path): 임시 템플릿 경로와 출력 경로

    Raises:
        FileNotFoundError: 템플릿 파일이 없는 경우

    Example:
        temp_template, temp_output = prepare_template(PO_TEMPLATE_FILE, "po")
        try:
            # ... 작업 수행 ...
        finally:
            cleanup_temp_file(temp_template)
    """
    if not template_path.exists():
        raise FileNotFoundError(f"템플릿 파일이 없습니다: {template_path}")

    timestamp = datetime.now().strftime('%Y%m%d%H%M%S')
    temp_dir = Path(tempfile.gettempdir())

    temp_template = temp_dir / f"{prefix}_template_{timestamp}.xlsx"
    temp_output = temp_dir / f"{prefix}_output_{timestamp}.xlsx"

    shutil.copy(template_path, temp_template)
    logger.debug(f"템플릿 복사 완료: {template_path} -> {temp_template}")

    return temp_template, temp_output


def cleanup_temp_file(temp_file: Path) -> None:
    """임시 파일을 안전하게 삭제

    Args:
        temp_file: 삭제할 임시 파일 경로
    """
    try:
        if temp_file.exists():
            temp_file.unlink()
            logger.debug(f"임시 파일 삭제: {temp_file}")
    except Exception as e:
        logger.warning(f"임시 파일 삭제 실패: {temp_file} - {e}")


# === 헤더 라벨 프리셋 ===

# Purchase Order 헤더 라벨 (PO, 이력 추출용)
PO_HEADER_LABELS: tuple[str, ...] = (
    'No.',
    'Item Number',
    'Item\nNumber',
    'Item',
)

# 거래명세표 헤더 라벨
TS_HEADER_LABELS: tuple[str, ...] = (
    '월/일',
    '품명',
    'DESCRIPTION',
)

# Proforma Invoice 헤더 라벨
PI_HEADER_LABELS: tuple[str, ...] = (
    'No.',
    'Description',
    'DESCRIPTION',
)

# Final Invoice 헤더 라벨
FI_HEADER_LABELS: tuple[str, ...] = (
    'No.',
    'Description',
    'DESCRIPTION',
)

# 기본 헤더 라벨 (모든 문서 유형에 공통)
DEFAULT_HEADER_LABELS: tuple[str, ...] = (
    'No.',
    'Item Number',
    'Item\nNumber',
    '품명',
    'Item',
)


def find_item_start_row_openpyxl(
    ws: OpenpyxlWorksheet,
    search_labels: tuple[str, ...] = DEFAULT_HEADER_LABELS,
    max_search_rows: int = 30,
    max_search_cols: int = 9,
    fallback_row: int = ITEM_START_ROW_FALLBACK,
) -> int:
    """템플릿에서 아이템 시작 행을 동적으로 찾기 (openpyxl 버전)

    헤더 레이블을 찾아서 그 다음 행이 아이템 시작 위치입니다.

    Args:
        ws: openpyxl Worksheet 객체
        search_labels: 검색할 헤더 레이블
        max_search_rows: 최대 검색 행 수
        max_search_cols: 최대 검색 열 수 (기본: 9, A-I)
        fallback_row: 헤더를 찾지 못했을 때 기본값

    Returns:
        아이템 시작 행 번호
    """
    for row in range(1, max_search_rows + 1):
        for col in range(1, max_search_cols + 1):
            cell_value = ws.cell(row=row, column=col).value
            if cell_value and any(
                label in str(cell_value) for label in search_labels
            ):
                logger.debug(
                    f"헤더 발견: Row {row}, 값='{cell_value}' -> 아이템 시작 Row {row + 1}"
                )
                return row + 1  # 레이블 다음 행이 데이터 시작

    logger.debug(f"헤더를 찾지 못함 -> 기본값 Row {fallback_row} 사용")
    return fallback_row


def find_item_start_row_xlwings(
    ws: xw.Sheet,
    search_labels: tuple[str, ...] = DEFAULT_HEADER_LABELS,
    max_search_rows: int = 30,
    columns: tuple[str, ...] = ('A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I'),
    fallback_row: int = ITEM_START_ROW_FALLBACK,
) -> int:
    """템플릿에서 아이템 시작 행을 동적으로 찾기 (xlwings 버전)

    헤더 레이블을 찾아서 그 다음 행이 아이템 시작 위치입니다.

    Args:
        ws: xlwings Sheet 객체
        search_labels: 검색할 헤더 레이블
        max_search_rows: 최대 검색 행 수
        columns: 검색할 열 문자 튜플
        fallback_row: 헤더를 찾지 못했을 때 기본값

    Returns:
        아이템 시작 행 번호
    """
    for row in range(1, max_search_rows + 1):
        for col in columns:
            cell_value = ws.range(f'{col}{row}').value
            if cell_value and any(
                label in str(cell_value) for label in search_labels
            ):
                logger.debug(
                    f"헤더 발견: Row {row}, 값='{cell_value}' -> 아이템 시작 Row {row + 1}"
                )
                return row + 1

    logger.debug(f"헤더를 찾지 못함 -> 기본값 Row {fallback_row} 사용")
    return fallback_row


# === 배치 연산 헬퍼 함수 (성능 최적화) ===

def batch_write_rows(
    ws: xw.Sheet,
    start_cell: str,
    data_2d: list[list],
) -> None:
    """2D 리스트를 한 번에 쓰기 (xlwings)

    여러 셀에 데이터를 쓸 때 셀 단위 COM 호출 대신 범위 쓰기를 사용하여
    성능을 크게 개선합니다.

    Args:
        ws: xlwings Sheet 객체
        start_cell: 시작 셀 주소 (예: 'A10')
        data_2d: 2D 리스트 (각 행은 리스트, [[row1], [row2], ...])

    Example:
        # 50개 아이템 * 8열 = 400회 COM 호출 → 1회로 감소
        data = [[date, name, remark, "EA", qty, price, amount, tax] for item in items]
        batch_write_rows(ws, 'A10', data)
    """
    if not data_2d:
        return

    num_rows = len(data_2d)
    num_cols = len(data_2d[0]) if data_2d else 0

    if num_cols == 0:
        return

    # 시작 셀에서 열 문자와 행 번호 추출
    col = ''.join(c for c in start_cell if c.isalpha())
    row = int(''.join(c for c in start_cell if c.isdigit()))

    # 끝 열 계산 (Z 이후 AA+ 다중 문자 열 지원)
    from openpyxl.utils import get_column_letter, column_index_from_string
    end_col = get_column_letter(column_index_from_string(col) + num_cols - 1)
    end_row = row + num_rows - 1

    # 한 번에 쓰기
    ws.range(f'{col}{row}:{end_col}{end_row}').value = data_2d
    logger.debug(f"배치 쓰기 완료: {col}{row}:{end_col}{end_row} ({num_rows}행 x {num_cols}열)")


def batch_read_column(
    ws: xw.Sheet,
    col: str,
    start_row: int,
    end_row: int,
) -> list:
    """열의 값을 한 번에 읽기 (xlwings)

    라벨 검색 등에서 셀 단위 COM 호출 대신 범위 읽기를 사용하여
    성능을 크게 개선합니다.

    Args:
        ws: xlwings Sheet 객체
        col: 열 문자 (예: 'A')
        start_row: 시작 행 번호
        end_row: 끝 행 번호

    Returns:
        값 리스트 (None 포함 가능)

    Example:
        # 36회 COM 호출 → 1회로 감소
        values = batch_read_column(ws, 'A', 15, 50)
        for idx, val in enumerate(values):
            if val and 'PO No' in str(val):
                return 15 + idx
    """
    values = ws.range(f'{col}{start_row}:{col}{end_row}').value

    # 단일 셀인 경우 리스트로 변환
    if not isinstance(values, list):
        values = [values]

    return values


def delete_rows_range(
    ws: xw.Sheet,
    start_row: int,
    count: int,
) -> None:
    """연속 행을 한 번에 삭제 (xlwings)

    반복 삭제 대신 범위 삭제를 사용하여 성능을 개선합니다.

    Args:
        ws: xlwings Sheet 객체
        start_row: 삭제 시작 행 번호
        count: 삭제할 행 수

    Example:
        # 5회 COM 호출 → 1회로 감소
        delete_rows_range(ws, 10, 5)  # Row 10-14 삭제
    """
    if count <= 0:
        return

    end_row = start_row + count - 1
    # xlShiftUp: 삭제 후 아래 행이 위로 올라옴
    ws.range(f'{start_row}:{end_row}').api.Delete(Shift=XlConstants.xlShiftUp)
    logger.debug(f"범위 삭제 완료: Row {start_row}-{end_row} ({count}행)")


def find_text_in_column_batch(
    ws: xw.Sheet,
    col: str,
    search_text: str,
    start_row: int,
    end_row: int,
) -> int | None:
    """배치 읽기로 열에서 텍스트 찾기 (xlwings)

    셀 단위 검색 대신 범위 읽기 후 Python에서 검색하여 성능을 개선합니다.

    Args:
        ws: xlwings Sheet 객체
        col: 검색할 열 (예: 'A')
        search_text: 찾을 텍스트 (부분 일치)
        start_row: 검색 시작 행
        end_row: 검색 끝 행

    Returns:
        찾은 행 번호 또는 None

    Example:
        # 36회 COM 호출 → 1회로 감소
        row = find_text_in_column_batch(ws, 'A', 'PO No', 15, 50)
    """
    values = batch_read_column(ws, col, start_row, end_row)

    for idx, val in enumerate(values):
        if val and search_text in str(val):
            return start_row + idx

    return None


# === 인쇄 레이아웃 헬퍼 (병합 셀 행 높이 · 한 페이지 채우기) ===

# 아이템 행의 최소 높이 (템플릿 기본값). 측정값이 이보다 작아도 이 아래로는 줄이지 않는다.
MIN_ITEM_ROW_HEIGHT: float = 15.0

# 측정값에 더할 여유. 병합 셀은 좌우 안쪽 여백이 단일 셀보다 미세하게 커서 딱 맞추면
# 마지막 줄의 아랫부분이 잘려 보일 수 있다.
ROW_HEIGHT_PAD: float = 2.0

# 측정용 보조 열 — 인쇄영역(A:I) 밖이어야 한다. 문서 5종이 모두 A:J만 쓴다.
PROBE_COLUMN: str = 'Z'

# 품목명 칸이 걸쳐 있는 열 — OC·FI·PI·CI·PL 템플릿이 모두 A:D를 행 단위로 병합한다(실측).
# **여기 한 곳에만 둔다.** 특히 CI와 PL은 선적서류라 고객에게 늘 같이 첨부되어 나란히
# 놓고 읽히므로, 두 문서의 줄 높이 규칙이 갈리면 바로 눈에 띈다.
ITEM_NAME_MERGED_COLS: str = 'ABCD'

# 아이템 그리드 내부 가로선 색 — 템플릿 상단 규칙선(#BBBBBB thin)과 같은 톤.
# 검정 thin은 PDF로 나가면 0.96pt 실선이라(캘리브레이션 실측: hairline=0.12 /
# thin=0.96 / medium=1.92pt) 격자 전체가 유독 무겁게 보인다(2026-08-05 보고).
# 내부선만 회색으로 누르고 표 프레임(헤더밴드 하단·마지막 행 하단)은 검정을 유지한다.
# COM의 Border.Color는 BGR 정수지만 회색은 대칭이라 RGB와 값이 같다.
ITEM_GRID_INNER_COLOR: int = 0xBBBBBB


def ensure_row_merges(ws: xw.Sheet, start_row: int, end_row: int) -> None:
    """아이템 행마다 품목명 칸(A:D) 병합을 다시 보장한다

    **행을 복사해 삽입하면 일부 행이 병합을 잃는다.** 48아이템 OC를 만들어 보니
    삽입한 41행 중 6행(28·38·45·51·54·57)에서 A:D 병합이 사라져 있었다
    (2026-08-03 실측). 그 행은 품목명이 A열(폭 약 51pt) 하나에 갇혀 5~7줄로 흘러
    표가 들쭉날쭉해진다 — 행 높이 측정 이전에 **문서 자체가 깨진 상태**다.

    Excel이 삽입 때 병합을 어떻게 처리했든 상관없도록, 범위를 통째로 풀고 행 단위로
    다시 병합한다 (`Merge(Across:=True)`는 각 행을 따로 병합한다). COM 호출 2회면 된다.

    값을 채운 뒤에 불러도 안전하다 — 병합은 좌상단(A열) 값만 남기는데 품목명이 거기 있다.

    병합 범위는 `ITEM_NAME_MERGED_COLS` 고정이다 — 문서 5종이 같은 값을 쓰는 것이
    불변식이라 인자로 열어 두지 않는다 (CI·PL은 늘 같이 첨부되어 나란히 읽힌다).

    Args:
        ws: xlwings Sheet
        start_row: 첫 행
        end_row: 마지막 행
    """
    if end_row < start_row:
        return

    cols = ITEM_NAME_MERGED_COLS
    span = f'{cols[0]}{start_row}:{cols[-1]}{end_row}'
    rng = ws.range(span).api
    rng.UnMerge()
    rng.Merge(True)  # Across=True — 범위 전체가 아니라 행마다 따로 병합
    logger.debug(f"행 병합 재적용: {span}")


def insert_copied_rows(ws: xw.Sheet, at_row: int, count: int, source_row: int) -> None:
    """원본 행 서식의 빈 행을 한 번에 삽입 (COM 3회)

    행당 `Copy()`+`Insert()` 왕복(2N회) 대신, 빈 행 N개를 **한 번에** 삽입한 뒤 원본
    행을 대상 범위에 **타일 복사**한다 — `Copy(Destination)`은 대상이 원본보다 크면
    원본을 반복해 채우고, 클립보드 상태도 남기지 않는다. 복사본에는 원본 행의 값도
    실리므로 내용은 지운다 (값을 채우기 전에 부르는 경로든, 빈 행을 덧대는 경로든
    이 규약이면 안전하다).

    주의: 복사·삽입을 거친 행은 병합(A:D)을 잃을 수 있다 — 호출부는 이 뒤에
    `ensure_row_merges()`(또는 `layout_item_rows()`)로 병합을 다시 보장할 것.

    Args:
        ws: xlwings Sheet
        at_row: 삽입 위치 (이 행부터 아래로 밀림)
        count: 삽입할 행 수
        source_row: 서식을 가져올 행 (보통 첫 아이템 행)
    """
    if count <= 0:
        return

    end = at_row + count - 1
    ws.range(f'{at_row}:{end}').api.Insert(Shift=XlConstants.xlShiftDown)
    ws.range(f'{source_row}:{source_row}').api.Copy(ws.range(f'{at_row}:{end}').api)
    ws.range(f'{at_row}:{end}').clear_contents()
    logger.debug(f"행 삽입: Row {at_row}부터 {count}개 (원본 Row {source_row})")


def _measure_wrapped_heights(
    ws: xw.Sheet,
    start_row: int,
    end_row: int,
    texts: list[str],
    cols: str,
) -> list[float]:
    """줄바꿈 텍스트가 요구하는 행 높이를 보조 열로 잰다 (pt)

    병합 셀은 Excel 자동 맞춤이 측정에서 제외하므로 직접 잴 수 없다 — 인쇄영역
    밖 보조 열(`PROBE_COLUMN`)의 폭을 병합 폭과 맞추고 같은 텍스트를 넣어
    Excel에게 재게 한다. 측정 원리·포인트 보정의 근거는 `autofit_merged_rows`
    docstring에 있다 (그쪽 실측을 이 함수가 그대로 수행한다).

    부수효과: 측정에 `rows.autofit()`을 쓰므로 대상 행들이 **잰 높이로 바뀐 채
    남는다** — 호출부가 최종 높이를 반드시 되써야 한다. 보조 열은 원복한다.

    Args:
        ws: xlwings Sheet
        start_row: 첫 행
        end_row: 마지막 행
        texts: 각 행에서 잴 텍스트 (길이 = 행 수)
        cols: 대상 병합이 걸친 열들 (예: 'ABCD')

    Returns:
        각 행의 측정 높이 (pt, 여유 미포함)
    """
    num_rows = end_row - start_row + 1
    source = ws.range(f'{cols[0]}{start_row}')
    char_width = sum(ws.range(f'{col}{start_row}').column_width for col in cols)
    target_pt = ws.range(f'{cols[0]}{start_row}:{cols[-1]}{start_row}').api.Width

    probe_range = f'{PROBE_COLUMN}{start_row}:{PROBE_COLUMN}{end_row}'
    original_width = ws.range(f'{PROBE_COLUMN}{start_row}').column_width

    try:
        probe = ws.range(probe_range)
        probe.column_width = char_width
        # 문자폭을 그대로 주면 열 패딩만큼 좁으므로, 실제 폭을 재서 비율로 보정한다
        probe_pt = ws.range(f'{PROBE_COLUMN}{start_row}').api.Width
        if probe_pt > 0 and target_pt > 0:
            probe.column_width = char_width * target_pt / probe_pt

        probe.api.WrapText = True
        probe.api.Font.Name = source.api.Font.Name
        probe.api.Font.Size = source.api.Font.Size
        probe.value = [[text] for text in texts]

        ws.range(f'{start_row}:{end_row}').rows.autofit()

        # 다중 행 RowHeight는 전부 같으면 그 값, 다르면 None — 균일한 문서(짧은
        # 품목명뿐)는 COM 1회로 끝나고, 섞였을 때만 행별로 읽는다.
        uniform = ws.range(probe_range).api.RowHeight
        if isinstance(uniform, (int, float)):
            measured = [float(uniform)] * num_rows
        else:
            measured = [
                ws.range(f'{PROBE_COLUMN}{row}').api.RowHeight
                for row in range(start_row, end_row + 1)
            ]
        # 진단용 — 이 경로는 재현이 까다로워서(병합 유실·폭 보정) 실측값을 남긴다.
        # f-string 인자는 즉시 평가되므로 COM 호출이 딸린 줄은 반드시 가드 안에 둔다.
        if logger.isEnabledFor(logging.DEBUG):
            logger.debug(
                f"측정: 병합폭 {target_pt:.2f}pt / "
                f"보조열 {ws.range(f'{PROBE_COLUMN}{start_row}').api.Width:.2f}pt / "
                f"폰트 {source.api.Font.Name} {source.api.Font.Size} -> {measured}"
            )
    finally:
        # 보조 열은 반드시 원복 — 남으면 인쇄영역 밖이라 안 보이는 채로 파일에 실린다
        ws.range(probe_range).clear_contents()
        ws.range(f'{PROBE_COLUMN}{start_row}').column_width = original_width

    return measured


def autofit_merged_rows(
    ws: xw.Sheet,
    start_row: int,
    end_row: int,
    texts: list[str],
) -> list[float]:
    """병합 셀이 든 행의 높이를 내용에 맞게 키운다

    **왜 `rows.autofit()`을 쓸 수 없나.** Excel의 자동 맞춤은 병합 셀을 측정 대상에서
    제외한다. 품목명 칸(A:D 병합)에 105자를 넣고 `autofit()`을 부르면 높이가
    15pt → 12.75pt로 **오히려 줄면서 1줄로 잘린다** (2026-08-03 실측). 즉 긴 품목명이
    조용히 사라지는데, 생성은 성공으로 끝나므로 PDF를 열어보기 전엔 모른다.

    **어떻게 재나.** 폰트 폭을 코드로 추정하지 않는다 — 인쇄영역 밖 보조 열의 폭을
    병합 셀과 같게 맞추고 같은 텍스트를 넣은 뒤, Excel에게 그 행을 재게 한다.
    보조 열은 병합돼 있지 않으므로 자동 맞춤이 정상 동작하고, 폭이 같으므로 줄 수도 같다.

    **폭은 문자 단위가 아니라 포인트로 맞춘다.** Excel은 열마다 안쪽 여백을 따로 붙여서,
    4개 열을 병합하면 1개 열보다 3개분 더 넓다 — OC 템플릿 실측으로 A:D는 198.00pt인데
    같은 문자폭(40.66)을 준 단일 열은 186.75pt였다. 이 11.25pt 차이 때문에 36자짜리
    품목명이 실제로는 한 줄에 들어가는데도 두 줄로 재어져 행이 쓸데없이 높아졌다
    (2026-08-03 실측). 그래서 한 번 재본 뒤 포인트 비율로 보정한다.
    (측정 자체는 `_measure_wrapped_heights`가 수행 — 주소 블록도 같은 원리를 쓴다.)

    측정 뒤에는 높이를 **명시적으로 되쓴다**. 보조 열을 비우면 행이 자동 모드로 남아
    다시 한 줄로 줄어들기 때문이다. 쓰기는 같은 높이의 연속 구간을 묶어 구간당 1회 —
    행높이 분포는 보통 두어 종류라(실측 48행 문서에서 3종) 행 수만큼 쓰지 않는다.

    Args:
        ws: xlwings Sheet
        start_row: 첫 아이템 행
        end_row: 마지막 아이템 행
        texts: 각 행에 들어간 텍스트 (길이가 행 수와 같아야 한다)

    Returns:
        각 행에 적용된 높이 (pt) — 페이지 계산에 그대로 쓴다

    Raises:
        ValueError: texts 길이가 행 수와 다른 경우 — 어긋난 채 조용히 진행하면
            그 아래 모든 행이 엉뚱한 높이를 받는다
    """
    num_rows = end_row - start_row + 1
    if num_rows <= 0:
        return []
    if len(texts) != num_rows:
        raise ValueError(f"texts {len(texts)}개 != 행 {num_rows}개 ({start_row}~{end_row})")

    measured = _measure_wrapped_heights(
        ws, start_row, end_row, texts, ITEM_NAME_MERGED_COLS
    )
    heights = [max(MIN_ITEM_ROW_HEIGHT, m + ROW_HEIGHT_PAD) for m in measured]

    # 같은 높이의 연속 구간을 묶어 쓴다 (행 범위 RowHeight 대입은 전 행에 적용된다)
    run_start = 0
    for idx in range(1, num_rows + 1):
        if idx == num_rows or heights[idx] != heights[run_start]:
            ws.range(
                f'{start_row + run_start}:{start_row + idx - 1}'
            ).api.RowHeight = heights[run_start]
            run_start = idx

    logger.debug(
        f"병합 셀 행 높이 조정: Row {start_row}-{end_row} "
        f"-> 높이 {[round(h, 1) for h in heights]}"
    )
    return heights


def address_row_heights(
    current: list[float],
    left_needs: list[float],
    right_need: float,
    pad: float = ROW_HEIGHT_PAD,
) -> list[float]:
    """주소 블록 행 높이 계산 (COM과 분리된 순수 함수)

    왼쪽은 행마다 독립된 한 줄 병합(A13:E13 꼴), 오른쪽은 같은 행들을 세로로
    걸친 병합 하나(G13:I15 꼴) — 행 높이는 두 요구를 동시에 만족해야 한다.

    - 각 행은 최소한 현재 높이를 유지한다 (짧은 주소 문서는 모양이 안 변한다)
    - 왼쪽 줄이 접히면 그 행만 자란다
    - 오른쪽 블록이 행 합보다 크면 부족분을 전 행에 균등 분배한다
      (한 행에 몰면 왼쪽 줄들의 간격이 들쭉날쭉해진다)

    Args:
        current: 각 행의 현재 높이 (pt)
        left_needs: 왼쪽 각 줄의 측정 높이 (pt, 여유 미포함)
        right_need: 오른쪽 블록 전체의 측정 높이 (pt, 없으면 0)
        pad: 병합 셀 여유 (`ROW_HEIGHT_PAD` 참조)

    Returns:
        각 행의 최종 높이 (pt)
    """
    heights = [
        max(cur, need + pad)
        for cur, need in zip(current, left_needs, strict=True)
    ]
    if right_need > 0 and heights:
        deficit = (right_need + pad) - sum(heights)
        if deficit > 0:
            share = deficit / len(heights)
            heights = [h + share for h in heights]
    return heights


def layout_address_rows(
    ws: xw.Sheet,
    start_row: int,
    left_texts: list[str],
    right_text: str,
    left_cols: str = 'ABCDE',
    right_cols: str = 'GHI',
) -> None:
    """헤더 주소 블록 손질: 줄바꿈 켜기 + 행 높이 확보 (OC·FI 공통)

    주소 칸은 병합 셀이라 넘친 텍스트가 옆 칸으로 흐르지 않고 **병합 경계에서
    잘린다** (2026-08-05 실측: 74자 bill-to가 A:E 끝에서, 81자 납품 주소가
    G:I 끝에서 잘린 채 PDF로 나감). wrap을 켜는 것만으로는 부족하다 — 병합
    셀은 autofit이 먹지 않아 행 높이가 그대로면 둘째 줄이 세로로 숨는다.
    그래서 아이템 행과 같은 보조 열 측정으로 높이까지 확보한다.

    값은 여기서 쓰지 않는다 — 생성기가 이미 셀에 쓴 텍스트를 측정용으로 받을
    뿐이다. 주소가 전부 짧으면 행 높이가 그대로라 기존 문서와 모양이 같다.

    Args:
        ws: xlwings Sheet
        start_row: 주소 첫 행 (OC=13, FI=12)
        left_texts: 왼쪽 각 행의 텍스트 (빈 문자열 허용) — 행 수를 정한다
        right_text: 오른쪽 세로 병합 블록의 텍스트 (없으면 '')
        left_cols: 왼쪽 병합이 걸친 열들
        right_cols: 오른쪽 병합이 걸친 열들
    """
    if not left_texts:
        return
    end_row = start_row + len(left_texts) - 1
    rows = range(start_row, end_row + 1)

    # 현재 높이는 측정 전에 읽는다 — 측정이 autofit으로 행 높이를 흐트러뜨린다
    current = [float(ws.range(f'{r}:{r}').api.RowHeight) for r in rows]

    ws.range(f'{left_cols[0]}{start_row}:{left_cols[-1]}{end_row}').api.WrapText = True
    ws.range(f'{right_cols[0]}{start_row}:{right_cols[-1]}{end_row}').api.WrapText = True

    left_needs = _measure_wrapped_heights(ws, start_row, end_row, left_texts, left_cols)
    right_need = 0.0
    if right_text:
        right_need = _measure_wrapped_heights(
            ws, start_row, start_row, [right_text], right_cols
        )[0]

    heights = address_row_heights(current, left_needs, right_need)
    for row, height in zip(rows, heights):
        ws.range(f'{row}:{row}').api.RowHeight = height
    logger.debug(
        f"주소 행 높이: Row {start_row}-{end_row} "
        f"{[round(c, 1) for c in current]} -> {[round(h, 1) for h in heights]}"
    )


def layout_item_rows(ws: xw.Sheet, start_row: int, texts: list[str]) -> list[float]:
    """아이템 행 마무리 손질: 병합 보장 → 행 높이 교정 (문서 5종 공통)

    행 복사·삽입은 병합을 잃을 수 있고(`ensure_row_merges` docstring의 실측),
    병합 칸에는 autofit이 먹지 않아 직접 재야 한다(`autofit_merged_rows`).
    이 둘은 늘 세트고 순서도 고정이라 한 함수로 묶는다 — 여섯 번째 문서가
    한쪽만 부르는 실수를 원천 차단한다. 값을 채운 **뒤에** 부른다.

    Args:
        ws: xlwings Sheet
        start_row: 첫 아이템 행
        texts: 각 행의 품목명 — 행 수는 len(texts)로 정한다

    Returns:
        각 행에 적용된 높이 (pt)
    """
    if not texts:
        return []
    end_row = start_row + len(texts) - 1
    ensure_row_merges(ws, start_row, end_row)
    return autofit_merged_rows(ws, start_row, end_row, texts)


